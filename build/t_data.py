# -*- coding: utf-8 -*-
"""build/t_data.py — 배치 T 자료 층: B0 판 동결 · B1 로더 · 가용 늦춤 · 선견 점검 틀.

설계 원본: 설계 워크플로 산출 tbatch_research.json 의 final(스크래치 · 저장소 밖) — 카드 12장
  T01 T02 T03 T04 T05 T08 T12 T13 T15 T16 T17 T18. 이 파일은 설계를 다시 짓지 않는다.
  옮기는 것은 final.data_build_plan 의 B0 · B1(+ B2 의 벤치마크 이음 · B3 의 원자료 로더)과
  final.tests.availability · final.tests.lookahead_check 뿐이다. 상태 · 신호(B4)는 여기 없다.

사용자 결정(2026-09-26)
  D1 장기 검정(1926+ French · Shiller · FRED · Cboe)은 사이트에 싣지 않는 «기전 증거 층»(L-R · L-C)에서만 쓴다.
     사이트는 10년 규약(build/maxyears.py MAX_YEARS)을 그대로 둔다 — 이 모듈의 어떤 산출도 사이트 자료가 아니다.
  D2 S&P 지수 선물(±5% 오버레이) · 지수 옵션 · 커버드콜 ETF 는 슬리브에 허용, 교차자산 선물(채권 · 원자재 · 통화)은 불허
     → T02(관리선물 TSMOM)는 측정만이고 전방 후보가 되지 않는다. 자료는 L 층 측정용으로만 받는다.
  D3 AQR · Cboe · Wurgler 원자료와 그 파생 NAV 는 커밋 · 게시하지 않는다 — 실행 때 저장소 밖 캐시로 받고 해시만 적는다.
     French · Shiller · yfinance 원자료도 재게시하지 않는다(캐시만). 저장소에 들어가는 입력은 연방 통계(FRED/ALFRED)뿐이다
     → data/_tb_fred/. 명세(data/_tb_manifest.json)에는 URL · 받은 때 · SHA-256 · 행 수 · 처음/끝 날짜만 적는다.
  D4 배치 FWER 0.025 는 H_T0 하나에 — 판정은 러너의 몫이고 이 모듈과 무관하다.

🚨 랩 규율 — 등록 커밋 전에는 수익 · 초과수익 · IR · 신호-수익 통계를 하나도 계산하거나 찍지 않는다.
   이 모듈이 찍는 것: 행 수 · 기간 · 결측 수 · 기업 수 미달 달 수 · 해시 · 판 표식 · 판 사이 개정 칸 수 ·
   (설계가 이미 적은) Mkt-RF 스폿 값의 일치 개수뿐이다. 다리 · 벤치마크 수익을 짓는 함수(french_vw_combine ·
   sp500_tr_splice)는 합성 자료 시험과, 실자료에서는 눈가린 연기 시험(--smoke: 표준출력 버림 · 모양만)으로만 돈다.

캐시 — 저장소 밖($TBATCH_CACHE, 없으면 %TEMP%/tbatch_cache). 저장소 안이면 멈춘다(D3).
  raw/<출처>/<파일>        고정본(명세의 SHA 가 가리키는 판) — 로더는 이것만 읽는다
  live/<YYYYMMDD>/<출처>/  --live 로 받은 지금 판(고정본과 대조만 하고 로더는 안 읽는다)
  meta/<키>.json           받은 때 · Last-Modified · 출처(설계 단계 사본인지 새로 받았는지)
  pylib/                   xlrd(Shiller .xls) — 이 PC 파이썬에 없으면 여기서 찾는다(pip install xlrd==2.0.2)

  python build/t_data.py --selftest          합성 자료 시험(망 · 캐시 · 저장소를 건드리지 않는다)
  python build/t_data.py --seed DIR          설계 단계 사본(DIR = 설계 워크플로의 tbatch 스크래치)을 고정본으로(SHA 대조)
  python build/t_data.py --fetch             고정본이 없는 자료만 받는다(yfinance · ALFRED 최초 공표 포함)
  python build/t_data.py --live              지금 판을 따로 받아 고정본과 대조한다(판 바뀜 보고 · 로더는 안 바뀐다)
  python build/t_data.py --manifest          고정본을 읽어 data/_tb_manifest.json 을 쓰고 범위를 보고한다
  python build/t_data.py --check             고정본 SHA == 명세 · 캐시 표지 · pylib xlrd 파일 SHA(실행 전 관문 — 틀리면 1)
  python build/t_data.py --cache-id-init     캐시 표지(meta/cache_id.txt · 무작위)를 만든다(없을 때만) — 명세가 그 SHA 를 적는다
  python build/t_data.py --backup [DIR]      고정본 · 표지 · pylib 를 저장소 밖 · %TEMP% 밖 사적 폴더로 SHA 대조 복사
                                             (기본 %LOCALAPPDATA%/yeodoo_private/tbatch_cache_backup)
  python build/t_data.py --restore [DIR]     백업 → 캐시(명세 SHA 와 같은 것만 · 다른 파일은 덮어쓰지 않는다)
  python build/t_data.py --smoke             실자료 눈가린 연기 시험(다리 · 이음 · 선견 틀 — 값은 안 본다)
  python build/t_data.py --b0 [PATH]         등록 문서 B0 절(해시 · French 절 머리글 · Last-Modified) 마크다운
"""
from __future__ import annotations

import csv
import datetime as _dt
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile

import numpy as np
import pandas as pd

try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
MANIFEST = os.path.join(DATA, "_tb_manifest.json")
PUB = os.path.join(DATA, "_tb_fred")                       # 연방 통계만 — 커밋 가능한 입력
CACHE = os.environ.get("TBATCH_CACHE") or os.path.join(tempfile.gettempdir(), "tbatch_cache")
PUBLIC_SRC = ("fred", "alfred")                            # 저장소에 고정본을 두는 출처(D3)
LOOKAHEAD_SEED = 20260925                                  # 설계의 씨앗(블록 셔플과 같은 값)
LOOKAHEAD_N = 200                                          # final.tests.lookahead_check «무작위 t 200개»
MIN_FIRMS = 20                                             # final.tests.F0 «다리 포트폴리오 기업 수 ≥ 20»

UA_BROWSER = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36"
UA_PY = "Python-urllib/3.12"

LIC = {
    "french": "French: 라이선스 문구 없음 — 출처 표기 · 원자료 재게시 안 함(캐시만)",
    "shiller": "Shiller: 문구 없음 — 원자료 재게시 안 함(캐시만)",
    "fred": "FRED/ALFRED 연방 통계 — 저장소 data/_tb_fred/ 에 고정",
    "aqr": "AQR: express prior written consent 없이 재배포·게시 금지(D3) — 캐시만 · 해시만 · 파생 결과 비공개",
    "cboe": "Cboe: 개인·비상업 사용(D3) — 캐시만 · 해시만 · 파생 NAV 비공개",
    "wurgler": "Wurgler: 라이선스 문구 없음(D3) — 커밋 안 함 · 실행 때 받는다",
    "yf": "yfinance: 개인 사용 — 캐시만",
}


# ══════════════════════════════════════════════════════════════════════════
#  자료 목록(설계 final.data · data.hashes · verified_this_step 를 옮긴 것)
# ══════════════════════════════════════════════════════════════════════════
FF = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/"
FF_VINTAGES = ("2005-08", "2015-08", "2020-08", "2024-08", "2024-12")   # F0 판별 FAST 불일치 · G5f(2024-12 FIZ · 2005-08)


def _ff_url(fname, vint=None):
    if vint is None:
        return FF + "ftp/" + fname
    if vint == "2024-12":
        return FF + "ftp_202412/" + fname                   # «FIZ 형식 마지막 판» 전 자료실 사본
    return FF + urllib.parse.quote("Data_Library/Historical_Archives/08 %s Update/ftp/%s" % (vint[:4], fname))


# id, 파일, 설계 SHA-256(verified_this_step · data.hashes), 설계가 적은 주 절(행 · 처음 · 끝), 쓰는 카드, 설계 단계 사본
FRENCH = [
    ("ff3_m", "F-F_Research_Data_Factors_CSV.zip", "b840dba55d319f4818fc7300e65c52eff5f64870c8d495fa58ff5d4cd749f5eb",
     (1201, "1926-07", "2026-07"), "B(Mkt TR) · RF · T01 · T04 · T05(IB) · T12 · T13 · 전 카드", "data/raw/french/F-F_Research_Data_Factors_CSV.zip"),
    ("ff3_d", "F-F_Research_Data_Factors_daily_CSV.zip", "1916d331c2c51d2aee3d00215897d2b8e5995cb387f1f4569ba46bff5fb049a8",
     (26296, "1926-07-01", "2026-07-31"), "T05 PANIC σ²(126거래일)", "data/raw/french/F-F_Research_Data_Factors_daily_CSV.zip"),
    ("mom_d", "F-F_Momentum_Factor_daily_CSV.zip", "b039d986db27fc831bcf1f52ba2134916e96fbeba4762fefc942429b0a568a96",
     None, "T05 BSC 쌍둥이(σ̂_UMD 126일)", "data/raw/french/F-F_Momentum_Factor_daily_CSV.zip"),
    ("prior12_2", "10_Portfolios_Prior_12_2_CSV.zip", "fb968db8061acb1052e595b59118f26fccda24963bd71f490a9a51c31eead8ec",
     (1195, "1927-01", "2026-07"), "T05 MOM(Hi PRIOR) · T04 쌍둥이(Lo PRIOR)", "data/raw/french/10_Portfolios_Prior_12_2_CSV.zip"),
    ("beta", "Portfolios_Formed_on_BETA_CSV.zip", "273ba6f791cc1b786d95e50c6741fee5b9daf4e598987596dd501cd07b45f890",
     (757, "1963-07", "2026-07"), "T03 L(Lo 20) · T04 DEF 쌍둥이 · T15 BETA 판", "data/raw/french/Portfolios_Formed_on_BETA_CSV.zip"),
    ("var", "Portfolios_Formed_on_VAR_CSV.zip", "c55289c24e14925bb3b3a2cfd7cc7cd4aedc8b7894fa58eaafa482cf757d6179",
     (757, "1963-07", "2026-07"), "T15 Lo/Hi 20 · T03 대조(VAR Lo 20)", "switch/data/Portfolios_Formed_on_VAR_CSV.zip"),
    ("op", "Portfolios_Formed_on_OP_CSV.zip", "43af7d9041520121b58ea85971aac084b29d16362840d67667d54b82eecef54b",
     (757, "1963-07", "2026-07"), "T04 DEF · T12 DEF(Hi 30)", "switch/data/Portfolios_Formed_on_OP_CSV.zip"),
    ("beme", "Portfolios_Formed_on_BE-ME_CSV.zip", "7ddd8918eb34b125b03e3e6e0bc927aef71dd9e3ff5a2a274d47da3a453c305b",
     (1201, "1926-07", "2026-07"), "T04 REB(Hi 30) · T16 · T17 십분위 반쪽 · T05 VAL 쌍둥이", "data/raw/french/Portfolios_Formed_on_BE-ME_CSV.zip"),
    ("ep", "Portfolios_Formed_on_E-P_CSV.zip", "5a235d4cde9fc59d9a46dd48724859fde2fb25586643fb990a4a2296d537bb23",
     (901, "1951-07", "2026-07"), "T16(E-P Hi 30)", "final/data/Portfolios_Formed_on_E-P_CSV.zip"),
    ("cfp", "Portfolios_Formed_on_CF-P_CSV.zip", "8dc6f9ba49ef7ab009122c0677113629ca7aa492222bb5911893e284aa930d61",
     (901, "1951-07", "2026-07"), "T16(CF-P Hi 30)", "final/data/Portfolios_Formed_on_CF-P_CSV.zip"),
    ("dp", "Portfolios_Formed_on_D-P_CSV.zip", "5fc9980adfe68c1e998431fe393b9bc762bf235cc5caec7d0c43f1fecd86b922",
     (1189, "1927-07", "2026-07"), "T16(D-P Hi 30)", "data/raw/french/Portfolios_Formed_on_D-P_CSV.zip"),
    ("ni", "Portfolios_Formed_on_NI_CSV.zip", "0923757ad94ac3e77506cb6b5700803370f8415f5a1ae2c99fd59fc9c9c88abb",
     (757, "1963-07", "2026-07"), "T18(< 0)", "final/data/Portfolios_Formed_on_NI_CSV.zip"),
    ("ind12", "12_Industry_Portfolios_CSV.zip", "c7316c6ae07bc2028632b57ad757dfef8303a62a6956946d7aabe926aadfdeb4",
     (1201, "1926-07", "2026-07"), "T04 · T12 방어 산업 쌍둥이(NoDur · Utils · Hlth)", "switch/data/12_Industry_Portfolios_CSV.zip"),
]
# 판(빈티지) — FF3 는 다섯 판(F0 · G5f), 나머지 French 파일은 2024-12 FIZ 판(G5f 재실행)
FRENCH_VINT = {
    ("ff3_m", "2005-08"): ("F-F_Research_Data_Factors_TXT.zip", "5788b202778d8249dfdc279590374f81f19432cd7576932d6df6ef8131a2c74a", "data/raw/french/vint_2005-08_TXT.zip"),
    ("ff3_m", "2015-08"): ("F-F_Research_Data_Factors_CSV.zip", None, "data/raw/french/vint_2015-08.zip"),
    ("ff3_m", "2020-08"): ("F-F_Research_Data_Factors_CSV.zip", None, "data/raw/french/vint_2020-08.zip"),
    ("ff3_m", "2024-08"): ("F-F_Research_Data_Factors_CSV.zip", None, "data/raw/french/vint_2024-08.zip"),
    ("ff3_m", "2024-12"): ("F-F_Research_Data_Factors_CSV.zip", "fc9e41cc3b66c62fa8f565a01fb4c8c0d7962c777303c46b8ab7102190b98ea2", "data/raw/french/vint_2024-12_FIZlast.zip"),
}
for _fid, _fn, _sha, _exp, _cards, _seed in FRENCH:
    if _fid != "ff3_m":
        FRENCH_VINT[(_fid, "2024-12")] = (_fn, None, None)

# 다리로 쓰는 칸(기업 수 ≥ 20 마스크 · F0) — 카드 규칙 · 쌍둥이에서 읽은 것
LEGS = {
    "prior12_2": ["Hi PRIOR", "Lo PRIOR"], "beta": ["Lo 20", "Hi 20"], "var": ["Lo 20", "Hi 20"], "op": ["Hi 30"],
    "beme": ["Hi 30", "Lo 10", "Dec 2", "Dec 3", "Dec 4", "Dec 5", "Dec 6", "Dec 7", "Dec 8", "Dec 9", "Hi 10"],
    "ep": ["Hi 30"], "cfp": ["Hi 30"], "dp": ["Hi 30"], "ni": ["< 0", "Lo 20"], "ind12": ["NoDur", "Utils", "Hlth"],
}
BEME_VALUE = ["Dec 6", "Dec 7", "Dec 8", "Dec 9", "Hi 10"]          # T17 가치 반쪽(6~10분위)
BEME_GROWTH = ["Lo 10", "Dec 2", "Dec 3", "Dec 4", "Dec 5"]         # T17 성장 반쪽(1~5분위)

# 설계가 이미 적은 F-F Mkt-RF 스폿(verified_this_step.FF3.spot_MktRF + T01 자료 칸 + B1) — 일치 개수만 찍는다
SPOT_MKTRF = {"1987-09": -2.58, "2020-01": -0.11, "2000-08": 7.03, "2007-10": 1.79, "2018-09": 0.06, "2021-12": 3.23,
              "2025-01": 2.80, "1932-06": -0.77, "1932-07": 33.61, "1974-09": -11.79, "1974-10": 16.11, "2009-02": -10.14,
              "2009-03": 9.01, "2018-12": -9.55, "2019-01": 8.37, "2020-03": -13.37, "2020-04": 13.60, "2022-06": -8.40,
              "2022-07": 9.57, "2022-09": -9.34, "2022-10": 7.85, "2025-04": -0.84, "2025-05": 6.06,
              "1929-10": -20.07, "1987-10": -23.19}

FRED = [  # id, 설계 SHA, 설계 기대(행 · 처음), 카드
    ("GS10", "e22128f6caa50e4e7fed03a1c2d323a517bad452a2c69eab89277450fb59bc5c", (881, "1953-04"), "T17 ΔGS10 · T17 실질금리 대리"),
    ("TB3MS", "ebf04b1ae5bc5729ba3bb2b35dbb0e0c5e22dace36bea0c02632782d625ae6ab", (1112, "1934-01"), "가용 목록(설계 availability)"),
    ("UNRATE", "ffe86c903f6944ebc8281e5474a15fc4ceaac61059bc51ec86a7241391d8713d", (944, "1948-01"), "T12(현재 판 — 최초 공표판의 대조)"),
    ("CPIAUCSL", "f8ecddf53a9a9a74dda92c2c4204e6039466fcb744bc74171b73b5f6aac48119", (956, "1947-01"), "T17 실질금리 대리(12개월 CPI)"),
    ("DFII10", None, None, "T17 S 층 DFII10 원판(국면 일치 수만)"),
]
ALFRED_SINGLE = ("UNRATE", "2008-12-05", "8b7a68367104b31926d1d397f73bc9cbf50a1529d7182dee10a72e75c63bf86d")
ALFRED_START = "1960-01"                                   # 최초 공표 계열의 첫 표본 판(월말) — T12 창 1964-07 에 앞선다

SHILLER_URL = ("https://img1.wsimg.com/blobby/go/e5e77e0b-59d1-44d9-ab25-4763ac982e53/downloads/"
               "70fec4f5-727f-4e53-b5f1-179af109c5fa/ie_data.xls?ver=1788371540009")
SHILLER_INDEX = "https://shillerdata.com/"
SHILLER = ("ie_data.xls", "044196dafe44c3030b2facbdea023975b3f6aa68b4e52f8f9bafc403e19589c1", (1869, "1871-01", "2026-09"),
           "T13 CAPE 재계산(P · E · CPI) · S&P TR 이음 1975-1987(D)", "data/raw/shiller/ie_data_shillerdata.xls")

CBOE = [  # 이름, 설계 SHA, 기대(행 · 처음 · 끝), 카드
    ("BXMD", "498886e4a22f707606cf5d3dbbe505dc059fafe92e6afb7dc14c2974d28aac79", (10136, "1986-06-20", "2026-09-22"), "T08 하락 국면 BXMD"),
    ("BXM", None, None, "T08 대조(BXM 베타 · 팩트시트 β 0.62 는 고정값)"),
    ("PUT", None, None, "옵션 가족 기록(측정 보조)"),
    ("VIX", "c1fa255f4e659d217f8cf244310a126517dc8f832568afe6be8184f46ca6dae7", (9278, "1990-01-02", None), "T08-S1 VRP(VIX 월말)"),
    ("SPX", "6bceb0b211f0841001ce6506a654e9b0bd27c90f069c8036f0d3a2930954c0fa", (13039, "1975-01-02", None), "S&P TR 이음 1975-1987"),
]
AQR = [  # id, 파일, 시트, 설계 SHA, 기대, 카드
    ("TSMOM", "Time-Series-Momentum-Factors-Monthly.xlsx", "TSMOM Factors",
     "33470930e2269c0d97be4732ec2d9c27ddbc69ac8133b059a263e27400263eeb", (497, "1985-01", "2026-05"), "T02(측정만 · D2 로 전방 불가)"),
    ("BAB", "Betting-Against-Beta-Equity-Factors-Monthly.xlsx", "BAB Factors", None, None, "T03 대조(AQR BAB · 측정만 · 라이선스 제한)"),
]
AQR_URL = "https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/"
WURGLER = ("SENTIMENT.xlsx", "https://pages.stern.nyu.edu/~jwurgler/data/SENTIMENT.xlsx",
           # 설계의 «DATA 817행» 은 머리글 줄을 센 것이다 — 1958-01~2025-12 는 816개월
           "3580398aaba6b72ad988844416c7b0ac9caca24fd2ab1d9149a983ec63b9f7a2", (816, "1958-01", "2025-12"), "T15 SENT^PIT(다섯 구성요소)",
           "critic/SENTIMENT.xlsx")
YF = [  # 티커, 설계가 적은 첫날(data.availability), 쓰임
    ("^GSPC", "1927-12-30", "사건 다리(지그재그) · T08-S1 실현분산 · 가격 상태 S 층"),
    ("^SP500TR", "1988-01-04", "S&P TR 이음 1988+ · G5g"),
    ("^VIX", "1990-01-02", "T08-S1 VRP 대조(Cboe VIX 와 같은 계열)"),
    ("SPY", "1993-01", "S 층 B"), ("BIL", "2007-05", "S 층 RF"),
    ("SPLV", "2011-05", "T03 · T15 ETF 쌍둥이"), ("SPHB", "2011-05", "T15 ETF 쌍둥이"), ("QUAL", "2013-07", "T04 · T12 DEF 쌍둥이"),
    ("RPV", "2006-03", "T04 REB · T16 쌍둥이"), ("IVE", "2000-05", "T04 · T16 · T17 대용"), ("IVW", "2000-05", "T17 대용"),
    ("IWD", "2000-05", "T17 쌍둥이"), ("IWF", "2000-05", "T17 쌍둥이"), ("VTV", None, "T17 쌍둥이"), ("VUG", None, "T17 쌍둥이"),
    ("MTUM", "2013-04", "T05 쌍둥이"), ("SPMO", "2015-10", "T05 쌍둥이"), ("XYLD", "2013-06", "T08 ETF 쌍둥이"),
    ("AQMIX", "2010-01-05", "T02 주 대용(측정만)"), ("DBMF", "2019-05", "T02 쌍둥이(측정만)"), ("RYMFX", "2007-02", "T02 쌍둥이(측정만)"),
    ("PKW", "2006-12-20", "T18 쌍둥이"), ("SYLD", "2013-05-14", "T18 쌍둥이"), ("COWZ", "2016-12-22", "T16 쌍둥이"),
]


def _key_safe(key):
    return re.sub(r"[^A-Za-z0-9_.@-]+", "_", key)


def registry():
    """모든 자료 한 벌 — 키 → 명세 칸. 키 꼴: «출처[@판]/id»."""
    R = {}

    def add(key, **kw):
        kw.setdefault("design_sha", None)
        kw.setdefault("expect", None)
        kw.setdefault("seed", None)
        kw.setdefault("ua", UA_BROWSER)
        kw.setdefault("vintage", None)
        kw["key"] = key
        R[key] = kw

    for fid, fn, sha, exp, cards, seed in FRENCH:
        add("french/" + fid, src="french", fid=fid, fname=fn, url=_ff_url(fn), license=LIC["french"], cards=cards,
            design_sha=sha, expect=exp, seed=seed, kind="french")
    cards_of = {f[0]: f[4] for f in FRENCH}
    for (fid, vint), (fn, sha, seed) in sorted(FRENCH_VINT.items()):
        add("french@%s/%s" % (vint, fid), src="french@" + vint, fid=fid, fname=fn, url=_ff_url(fn, vint), license=LIC["french"],
            cards=("F0 판별 FAST 불일치 · G5f" if fid == "ff3_m" else "G5f(2024-12 판 재실행) · ") + ("" if fid == "ff3_m" else cards_of[fid]),
            design_sha=sha, seed=seed, kind="french", vintage=vint)
    for sid, sha, exp, cards in FRED:
        add("fred/" + sid, src="fred", fid=sid, fname=sid + ".csv", url="https://fred.stlouisfed.org/graph/fredgraph.csv?id=" + sid,
            license=LIC["fred"], cards=cards, design_sha=sha, expect=exp, seed="data/raw/fred/%s.csv" % sid, ua=UA_PY, kind="fred")
    s, vd, sha = ALFRED_SINGLE
    add("fred/alfred_%s_%s" % (s, vd), src="fred", fid="alfred_%s_%s" % (s, vd), fname="alfred_%s_%s.csv" % (s, vd),
        url="https://alfred.stlouisfed.org/graph/alfredgraph.csv?id=%s&vintage_date=%s" % (s, vd), license=LIC["fred"],
        cards="T12 개정 한계 대조(설계: 2008 판 대비 731개월 중 15개월 차이)", design_sha=sha, seed="data/raw/fred/alfred_%s_%s.csv" % (s, vd),
        ua=UA_PY, kind="fred")
    add("alfred/UNRATE_first", src="alfred", fid="UNRATE_first", fname="UNRATE_first_release.csv",
        url="https://alfred.stlouisfed.org/graph/alfredgraph.csv?id=UNRATE,…&vintage_date=<월말 12개씩>", license=LIC["fred"],
        cards="T12 최초 공표 UNRATE(t−1 · 최초 공표 판 날짜 ≤ t)", ua=UA_PY, kind="alfred_first")
    fn, sha, exp, cards, seed = SHILLER
    add("shiller/ie_data", src="shiller", fid="ie_data", fname=fn, url=SHILLER_URL, license=LIC["shiller"], cards=cards,
        design_sha=sha, expect=exp, seed=seed, kind="shiller")
    for nm, sha, exp, cards in CBOE:
        add("cboe/" + nm, src="cboe", fid=nm, fname=nm + "_History.csv",
            url="https://cdn.cboe.com/api/global/us_indices/daily_prices/%s_History.csv" % nm, license=LIC["cboe"], cards=cards,
            design_sha=sha, expect=exp, seed="data/raw/cboe/%s_History.csv" % nm, ua=UA_PY, kind="cboe")
    for aid, fn, sheet, sha, exp, cards in AQR:
        add("aqr/" + aid, src="aqr", fid=aid, fname=fn, url=AQR_URL + fn, license=LIC["aqr"], cards=cards, design_sha=sha,
            expect=exp, seed="data/raw/aqr/" + fn, kind="aqr", sheet=sheet)
    fn, url, sha, exp, cards, seed = WURGLER
    add("wurgler/SENTIMENT", src="wurgler", fid="SENTIMENT", fname=fn, url=url, license=LIC["wurgler"], cards=cards,
        design_sha=sha, expect=exp, seed=seed, kind="wurgler")
    for tk, first, cards in YF:
        add("yf/" + tk, src="yf", fid=tk, fname=tk.replace("^", "_") + ".csv", url="yfinance:Ticker(%r).history(period='max', auto_adjust=False, actions=True)" % tk,
            license=LIC["yf"], cards=cards, expect=(None, first, None) if first else None, kind="yf")
    return R


# ══════════════════════════════════════════════════════════════════════════
#  캐시 · 해시 · 받기
# ══════════════════════════════════════════════════════════════════════════
def _cache_guard():
    c, r = os.path.abspath(CACHE), os.path.abspath(ROOT)
    if os.path.normcase(c).startswith(os.path.normcase(r) + os.sep) or os.path.normcase(c) == os.path.normcase(r):
        raise SystemExit("🚨 캐시(%s)가 저장소 안이다 — 사용 허락이 없는 원자료는 저장소 밖에만 둔다(D3)." % CACHE)
    return c


def sha256_bytes(b):
    return hashlib.sha256(b).hexdigest()


def _now():
    return _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def pinned_path(ds):
    if ds["src"] in PUBLIC_SRC:
        return os.path.join(PUB, ds["fname"])
    return os.path.join(_cache_guard(), "raw", ds["src"], ds["fname"])


def _meta_path(ds):
    return os.path.join(_cache_guard(), "meta", _key_safe(ds["key"]) + ".json")


def _read_meta(ds):
    p = _meta_path(ds)
    if os.path.exists(p):
        return json.load(io.open(p, encoding="utf-8"))
    return None


def _write_meta(ds, meta):
    p = _meta_path(ds)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    io.open(p, "w", encoding="utf-8").write(json.dumps(meta, ensure_ascii=False, indent=1))


def _write_bytes(path, blob):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".part"
    with open(tmp, "wb") as f:
        f.write(blob)
    os.replace(tmp, path)


def read_pinned(ds):
    p = pinned_path(ds)
    if not os.path.exists(p):
        raise FileNotFoundError("고정본 없음: %s — python build/t_data.py --seed DIR 또는 --fetch" % ds["key"])
    with open(p, "rb") as f:
        return f.read()


# ══════════════════════════════════════════════════════════════════════════
#  캐시 표지 · 파이썬 라이브러리 핀 · 백업(적대 검토 2026-09-26)
#  · 캐시 표지 — 캐시 뿌리마다 무작위 표지(meta/cache_id.txt)를 두고 그 SHA-256 을 명세에 적는다. 러너는 명세와 같은 표지를 가진
#    캐시에서만 굽는다 → TBATCH_CACHE 를 새 폴더로 돌려(시작 표식 · 산출이 없는 캐시) 몰래 다시 굽는 길을 막는다.
#  · pylib — Shiller .xls 를 읽는 xlrd 가 이 PC 파이썬에 없어 캐시 pylib/ 에서 읽는다. 그 .py 파일들의 SHA-256 을 명세에 적고
#    러너가 실제로 불러온 xlrd 의 파일과 대조한다(판 문자열만으로는 코드가 고정되지 않는다).
#  · 백업 — 고정본(AQR · Cboe · Shiller · Wurgler · French · yfinance)은 다시 받아도 같은 SHA 가 나오지 않는다(Cboe 는 이미 바뀌었고
#    AQR 는 갱신마다 전체 이력을 다시 짓는다). %TEMP% 는 정리될 수 있으므로 저장소 밖 · %TEMP% 밖 사적 폴더에 SHA 를 대조하며 복사한다.
#    백업도 저장소 밖이다(D3 — 커밋 · 게시하지 않는다).
# ══════════════════════════════════════════════════════════════════════════
CACHE_ID = os.path.join("meta", "cache_id.txt")
PYLIB_PKG = "xlrd"
BACKUP_DEFAULT = os.path.join(os.environ.get("LOCALAPPDATA") or os.path.expanduser("~"), "yeodoo_private", "tbatch_cache_backup")


def _inside(p, root):
    a, r = os.path.normcase(os.path.abspath(p)), os.path.normcase(os.path.abspath(root))
    return a == r or a.startswith(r + os.sep)


def cache_id_path():
    return os.path.join(_cache_guard(), CACHE_ID)


def cache_id_init():
    """캐시 표지가 없으면 만든다(무작위 32바이트 hex). 이미 있으면 그대로 — 돌려주는 것: 새로 만들었는가."""
    import secrets
    p = cache_id_path()
    if os.path.exists(p):
        return False
    os.makedirs(os.path.dirname(p), exist_ok=True)
    io.open(p, "w", encoding="utf-8", newline="\n").write(secrets.token_hex(32) + "\n")
    return True


def cache_id_sha():
    p = cache_id_path()
    if not os.path.exists(p):
        return None
    with open(p, "rb") as f:
        return sha256_bytes(f.read())


def check_cache_identity(doc=None):
    """캐시 표지 SHA == 명세 — (ok, 사유)."""
    doc = json.load(io.open(MANIFEST, encoding="utf-8")) if doc is None else doc
    want = (doc.get("cache_identity") or {}).get("sha256")
    have = cache_id_sha()
    if not want:
        return False, "명세에 캐시 표지가 없다"
    if have is None:
        return False, "캐시에 표지(%s)가 없다 — 등록한 캐시가 아니다(백업에서 --restore)" % CACHE_ID.replace(os.sep, "/")
    if have != want:
        return False, "캐시 표지가 명세와 다르다 — 등록한 캐시 뿌리가 아니다"
    return True, "캐시 표지 일치"


def pylib_pins():
    """지금 불러오는 xlrd 의 판과 .py 파일 SHA-256(파일 이름만 · 경로 없음)."""
    x = _xlrd()
    d = os.path.dirname(os.path.abspath(x.__file__))
    files = {}
    for fn in sorted(os.listdir(d)):
        if fn.endswith(".py"):
            with open(os.path.join(d, fn), "rb") as f:
                files[fn] = sha256_bytes(f.read())
    return {PYLIB_PKG: {"version": getattr(x, "__version__", None), "files": files,
                        "from_cache_pylib": _inside(d, os.path.join(_cache_guard(), "pylib"))}}


def check_pylib(doc=None):
    doc = json.load(io.open(MANIFEST, encoding="utf-8")) if doc is None else doc
    want = ((doc.get("pylib") or {}).get(PYLIB_PKG) or {})
    if not want:
        return False, "명세에 pylib 핀이 없다"
    have = pylib_pins()[PYLIB_PKG]
    if have["version"] != want.get("version"):
        return False, "xlrd 판 %s ≠ 명세 %s" % (have["version"], want.get("version"))
    if have["files"] != want.get("files"):
        diff = sorted(set(have["files"].items()) ^ set((want.get("files") or {}).items()))
        return False, "xlrd 파일이 명세와 다르다(%s …)" % ", ".join(sorted({k for k, _ in diff})[:3])
    return True, "xlrd %s 파일 %d개 일치" % (have["version"], len(have["files"]))


def _backup_items():
    """(캐시 안 상대 경로, 기대 SHA 또는 None) — 라이선스 · 비연방 고정본 전부 + 캐시 표지 + pylib xlrd."""
    doc = json.load(io.open(MANIFEST, encoding="utf-8"))
    out = []
    for key, ds in registry().items():
        if ds["src"] in PUBLIC_SRC:
            continue                                     # 연방 통계는 저장소가 고정한다
        rec = doc["datasets"].get(key) or {}
        out.append((os.path.join("raw", ds["src"], ds["fname"]), (rec.get("pinned") or {}).get("sha256")))
    out.append((CACHE_ID, (doc.get("cache_identity") or {}).get("sha256")))
    for fn, sha in (((doc.get("pylib") or {}).get(PYLIB_PKG) or {}).get("files") or {}).items():
        out.append((os.path.join("pylib", PYLIB_PKG, fn), sha))
    return out


def backup(dest=None):
    """고정본 · 표지 · pylib 를 dest 로 SHA 대조 복사(저장소 안이면 멈춘다 · %TEMP% 안이면 경고). 값을 읽지 않는다."""
    dest = os.path.abspath(dest or BACKUP_DEFAULT)
    if _inside(dest, ROOT):
        raise SystemExit("🚨 백업이 저장소 안이다(D3): %s" % dest)
    warn = _inside(dest, tempfile.gettempdir())
    src = _cache_guard()
    rows = []
    for rel, want in _backup_items():
        a = os.path.join(src, rel)
        if not os.path.exists(a):
            rows.append({"path": rel.replace(os.sep, "/"), "state": "캐시에 없음"})
            continue
        with open(a, "rb") as f:
            blob = f.read()
        sha = sha256_bytes(blob)
        if want and sha != want:
            raise SystemExit("🚨 캐시 %s 가 명세와 다르다 — 백업하지 않는다" % rel)
        b = os.path.join(dest, rel)
        if os.path.exists(b):
            with open(b, "rb") as f:
                if sha256_bytes(f.read()) == sha:
                    rows.append({"path": rel.replace(os.sep, "/"), "state": "이미 같음", "sha256": sha})
                    continue
        _write_bytes(b, blob)
        with open(b, "rb") as f:
            if sha256_bytes(f.read()) != sha:
                raise SystemExit("🚨 백업 사본 SHA 가 다르다: %s" % rel)
        rows.append({"path": rel.replace(os.sep, "/"), "state": "복사", "sha256": sha})
    io.open(os.path.join(dest, "BACKUP_INDEX.json"), "w", encoding="utf-8", newline="\n").write(
        json.dumps({"made_at": _now(), "items": rows, "note": "배치 T 고정본 백업(D3 — 저장소 밖 · 게시 금지 · 값 없음)"},
                   ensure_ascii=False, indent=1) + "\n")
    return {"dest_in_temp": warn, "n": len(rows), "copied": sum(1 for r in rows if r["state"] == "복사"),
            "same": sum(1 for r in rows if r["state"] == "이미 같음"), "missing": [r["path"] for r in rows if r["state"] == "캐시에 없음"]}


def restore(src_dir=None):
    """백업 → 캐시(명세 SHA 와 같은 것만 · 캐시에 이미 다른 파일이 있으면 멈춘다)."""
    src_dir = os.path.abspath(src_dir or BACKUP_DEFAULT)
    dst = _cache_guard()
    n = 0
    for rel, want in _backup_items():
        a = os.path.join(src_dir, rel)
        if not os.path.exists(a):
            raise SystemExit("🚨 백업에 %s 가 없다" % rel)
        with open(a, "rb") as f:
            blob = f.read()
        if want and sha256_bytes(blob) != want:
            raise SystemExit("🚨 백업 %s 가 명세와 다르다" % rel)
        b = os.path.join(dst, rel)
        if os.path.exists(b):
            with open(b, "rb") as f:
                if sha256_bytes(f.read()) != sha256_bytes(blob):
                    raise SystemExit("🚨 캐시에 다른 %s 가 이미 있다 — 덮어쓰지 않는다" % rel)
            continue
        _write_bytes(b, blob)
        n += 1
    return {"restored": n}


def _http_get(url, ua, timeout=90, tries=3):
    last = None
    for k in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": ua})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                blob = r.read()
                return blob, {"status": r.status, "last_modified": r.headers.get("Last-Modified"),
                              "content_type": r.headers.get("Content-Type")}
        except urllib.error.HTTPError as e:
            if e.code in (400, 403, 404, 410):
                raise
            last = e
        except Exception as e:                      # 망 오류 — 잠깐 쉬고 다시
            last = e
        time.sleep(2.0 * (k + 1))
    raise last


def _sniff(ds, blob):
    """받은 것이 기대한 꼴인가(오류 페이지를 고정본으로 삼지 않게)."""
    k = ds["kind"]
    if k in ("french", "aqr", "wurgler"):
        ok = blob[:2] == b"PK"
    elif k == "shiller":
        ok = blob[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    elif k == "cboe":
        ok = blob[:4].upper() == b"DATE"
    elif k == "fred":
        ok = blob[:16].lower().startswith(b"observation_date") or blob[:4] == b"DATE"
    else:
        ok = len(blob) > 0
    if not ok:
        raise ValueError("%s: 받은 내용이 기대한 꼴이 아니다(앞 %r)" % (ds["key"], blob[:40]))


def _shiller_live_url():
    """shillerdata.com 이 지금 가리키는 ie_data.xls 링크(없으면 고정 URL)."""
    try:
        html, _ = _http_get(SHILLER_INDEX, UA_BROWSER, timeout=60)
        m = re.search(rb"https://img1\.wsimg\.com/[^\"'\s<>]*ie_data\.xls[^\"'\s<>]*", html)
        if m:
            return m.group(0).decode("ascii", "replace").replace("&amp;", "&")
    except Exception:
        pass
    return SHILLER_URL


def _yf_blob(ticker):
    import yfinance as yf
    df = yf.Ticker(ticker).history(period="max", auto_adjust=False, actions=True)
    if df is None or df.empty:
        raise RuntimeError("yfinance %s: 빈 응답" % ticker)
    idx = pd.to_datetime(df.index)
    if idx.tz is not None:
        idx = idx.tz_localize(None)
    df.index = idx.normalize()
    df.index.name = "Date"
    cols = [c for c in ("Open", "High", "Low", "Close", "Adj Close", "Volume", "Dividends", "Stock Splits", "Capital Gains") if c in df.columns]
    txt = df[cols].to_csv(date_format="%Y-%m-%d", float_format="%.10g", lineterminator="\n")
    return txt.encode("utf-8"), {"status": 200, "last_modified": None, "yfinance": getattr(yf, "__version__", "?")}


def _download(ds):
    if ds["kind"] == "yf":
        return _yf_blob(ds["fid"]) + (ds["url"],)
    url = _shiller_live_url() if ds["kind"] == "shiller" else ds["url"]
    blob, info = _http_get(url, ds["ua"])
    _sniff(ds, blob)
    return blob, info, url


def fetch(ds, force=False, sleep=0.5):
    """고정본이 없으면 받아 고정본으로 둔다(있으면 손대지 않는다 — 판 동결)."""
    if ds["kind"] == "alfred_first":
        p = pinned_path(ds)
        return p if os.path.exists(p) and not force else alfred_first_fetch()
    p = pinned_path(ds)
    if os.path.exists(p) and not force:
        return p
    blob, info, url = _download(ds)
    _write_bytes(p, blob)
    _write_meta(ds, {"key": ds["key"], "url": url, "fetched_at": _now(), "last_modified": info.get("last_modified"),
                     "origin": "live", "sha256": sha256_bytes(blob), "bytes": len(blob),
                     **({"yfinance": info["yfinance"]} if "yfinance" in info else {})})
    time.sleep(sleep)
    return p


def fetch_live(ds, day=None, sleep=0.5):
    """지금 판을 live/<날짜>/ 에 받아 정보만 돌려준다(고정본 · 로더는 그대로)."""
    if ds["kind"] in ("yf", "alfred_first"):
        return None
    day = day or _dt.date.today().strftime("%Y%m%d")
    p = os.path.join(_cache_guard(), "live", day, ds["src"], ds["fname"])
    mp = p + ".meta.json"
    if os.path.exists(p) and os.path.exists(mp):
        return json.load(io.open(mp, encoding="utf-8")) | {"path": p}
    try:
        blob, info, url = _download(ds)
    except Exception as e:
        return {"error": repr(e)[:300], "fetched_at": _now()}
    _write_bytes(p, blob)
    meta = {"url": url, "fetched_at": _now(), "last_modified": info.get("last_modified"), "sha256": sha256_bytes(blob), "bytes": len(blob)}
    io.open(mp, "w", encoding="utf-8").write(json.dumps(meta, ensure_ascii=False, indent=1))
    time.sleep(sleep)
    return meta | {"path": p}


def seed(design_dir, R=None):
    """설계 단계 사본을 고정본으로 들인다 — 설계 SHA 가 있으면 같아야만 들인다(다르면 들이지 않고 보고)."""
    R = R or registry()
    out = {}
    for key, ds in R.items():
        if not ds.get("seed"):
            continue
        src = os.path.join(design_dir, ds["seed"])
        p = pinned_path(ds)
        if os.path.exists(p):
            out[key] = "이미 고정본 있음"
            continue
        if not os.path.exists(src):
            out[key] = "사본 없음"
            continue
        with open(src, "rb") as f:
            blob = f.read()
        sha = sha256_bytes(blob)
        if ds["design_sha"] and sha != ds["design_sha"]:
            out[key] = "🚨 사본 SHA 가 설계와 다르다 — 들이지 않음(%s…)" % sha[:12]
            continue
        _sniff(ds, blob)
        _write_bytes(p, blob)
        lm = None
        hdr = src + ".hdr.json"
        if os.path.exists(hdr):
            try:
                lm = json.load(io.open(hdr, encoding="utf-8")).get("last_modified")
            except Exception:
                lm = None
        mt = _dt.datetime.fromtimestamp(os.path.getmtime(src), _dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        _write_meta(ds, {"key": key, "url": ds["url"], "fetched_at": mt, "last_modified": lm,
                         "origin": "design-stage copy (설계 워크플로가 받은 판 · 2026-09-25)", "sha256": sha, "bytes": len(blob)})
        out[key] = "들임" + (" · 설계 SHA 일치" if ds["design_sha"] else " · 설계 SHA 없음(이 사본이 고정본)")
    return out


# ══════════════════════════════════════════════════════════════════════════
#  French 다절 CSV/TXT 파서
# ══════════════════════════════════════════════════════════════════════════
_DATE_LENS = (4, 6, 8)
MISSING_CODES = (-99.99, -999.0)
RETURN_KINDS = ("m", "d", "a", "vw_m", "ew_m", "vw_a", "ew_a")


def _is_data_line(ln):
    tok = ln.replace(",", " ").split()
    return bool(tok) and tok[0].isdigit() and len(tok[0]) in _DATE_LENS


def _norm_col(c):
    c = re.sub(r"\s+", " ", c.strip())
    m = re.match(r"^(\d)-Dec$", c)                 # BE-ME 파일의 «2-Dec» 는 «Dec 2»(NI · E-P 파일 표기)
    return "Dec %s" % m.group(1) if m else c


def _block_kind(title, dlen):
    t = title.lower().replace("value-weighted", "value weight").replace("weighted", "weight")
    ann = "annual" in t or dlen == 4
    if "number of firms" in t:
        return "nfirms"
    if "average firm size" in t:
        return "avgsize"
    if "value weight average of" in t:
        return "vwavg" + ("_a" if dlen == 4 else "")
    if t.startswith("sum of"):
        return "sum_ratio"
    if "equal weight" in t and "return" in t:
        return "ew_a" if ann else "ew_m"
    if "value weight" in t and "return" in t:
        return "vw_a" if ann else "vw_m"
    if t == "" or "factors" in t:
        return {4: "a", 6: "m", 8: "d"}[dlen]
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")[:40]


def _num(tok):
    tok = tok.strip()
    if tok == "":
        return np.nan
    v = float(tok)
    for m in MISSING_CODES:
        if abs(v - m) < 1e-9:
            return np.nan
    return v


def _period_index(dates, dlen):
    if dlen == 6:
        return pd.PeriodIndex(["%s-%s" % (d[:4], d[4:6]) for d in dates], freq="M")
    if dlen == 8:
        return pd.DatetimeIndex(pd.to_datetime(dates, format="%Y%m%d"))
    return pd.Index([int(d) for d in dates], name="year")


def parse_french_text(txt):
    """French 자료실 파일(CSV 여러 절 또는 옛 TXT) → {"blocks": {kind: DataFrame(원 단위)}, "titles", "cols", "crsp", "preamble", "n_missing"}.
    절 = 제목 줄(없을 수 있다) + 머리글 줄 + 날짜(4 · 6 · 8자리)로 시작하는 줄들. −99.99 · −999 는 NaN."""
    lines = txt.replace("\r", "").split("\n")
    raw_blocks, cur = [], None

    def _cols_of(h):
        return [c for c in h.split(",")][1:] if "," in h else h.split()
    for i, ln in enumerate(lines):
        if _is_data_line(ln):
            if cur is None:
                j = i - 1
                hdr = lines[j] if j >= 0 else ""
                if hdr.strip() == "":
                    # 머리글과 자료 사이에 빈 줄이 낀 절(현재 VAR 판의 VW 연간 절) — 가장 가까운 비지 않은 줄이 자료 칸 수와 맞는 머리글이면 그것을 쓴다.
                    #   맞지 않으면 예전처럼(머리글 없음 · c1…) 둔다(적대 검토 2026-09-26 — 조용히 틀린 제목 · 칸 이름을 막는다).
                    k = j
                    while k >= 0 and lines[k].strip() == "":
                        k -= 1
                    ntok = len(ln.split(",")) if "," in ln else len(ln.split())
                    if k >= 0 and not _is_data_line(lines[k]) and len(_cols_of(lines[k])) == ntok - 1 and ntok > 1:
                        j, hdr = k, lines[k]
                title = lines[j - 1].strip() if j - 1 >= 0 else ""
                cols = _cols_of(hdr)
                cur = {"title": title, "cols": [_norm_col(c) for c in cols], "rows": [], "hdr_line": j}
                raw_blocks.append(cur)
            toks = [t for t in ln.split(",")] if "," in ln else ln.split()
            cur["rows"].append(toks)
        else:
            cur = None
    if not raw_blocks:
        raise ValueError("French 파일에서 자료 절을 못 찾았다")
    first_hdr = raw_blocks[0]["hdr_line"]
    preamble = [ln.strip() for ln in lines[:max(0, first_hdr - (1 if raw_blocks[0]["title"] else 0))] if ln.strip()]
    m = re.search(r"(\d{6}) CRSP database", " ".join(preamble))
    out = {"blocks": {}, "titles": {}, "cols": {}, "n_missing": {}, "dup_dates": {}, "crsp": m.group(1) if m else None,
           "preamble": preamble}
    for b in raw_blocks:
        dlen = len(b["rows"][0][0].strip())
        kind = _block_kind(b["title"], dlen)
        k2, n = kind, 2
        while k2 in out["blocks"]:
            k2, n = "%s_%d" % (kind, n), n + 1
        cols = b["cols"] or ["c%d" % j for j in range(1, len(b["rows"][0]))]
        vals, dates, nmiss = [], [], 0
        for r in b["rows"]:
            d = r[0].strip()
            if len(d) != dlen:
                continue
            row = [_num(x) for x in r[1:1 + len(cols)]]
            row += [np.nan] * (len(cols) - len(row))
            nmiss += sum(1 for x, t in zip(row, r[1:1 + len(cols)]) if x != x and t.strip() != "")
            vals.append(row)
            dates.append(d)
        idx = _period_index(dates, dlen)
        df = pd.DataFrame(vals, index=idx, columns=cols)
        dup = int(df.index.duplicated().sum())
        df = df[~df.index.duplicated(keep="first")].sort_index()
        out["blocks"][k2] = df
        out["titles"][k2] = b["title"]
        out["cols"][k2] = cols
        out["n_missing"][k2] = nmiss
        out["dup_dates"][k2] = dup
    return out


def parse_french_blob(blob):
    z = zipfile.ZipFile(io.BytesIO(blob))
    names = [n for n in z.namelist() if not n.endswith("/")]
    txt = z.read(names[0]).decode("latin-1")
    out = parse_french_text(txt)
    out["inner"] = names[0].split("/")[-1]
    return out


_FR_CACHE = {}


def load_french(fid, vintage=None):
    key = ("french/%s" % fid) if vintage is None else ("french@%s/%s" % (vintage, fid))
    if key not in _FR_CACHE:
        R = registry()
        if key not in R:
            raise KeyError("French 자료 목록에 없다: %s" % key)
        _FR_CACHE[key] = parse_french_blob(read_pinned(R[key]))
    return _FR_CACHE[key]


def french(fid, kind=None, vintage=None, pct=True):
    """French 표 한 절. 수익 절(m · d · a · vw_* · ew_*)은 pct=True 면 소수(÷100). 기업 수 · 평균 규모 · 비율 절은 원 단위."""
    f = load_french(fid, vintage)
    B = f["blocks"]
    if kind is None:
        kind = next(k for k in ("vw_m", "m", "d") if k in B)
    df = B[kind].copy()
    if pct and re.sub(r"_\d+$", "", kind) in RETURN_KINDS:
        df = df / 100.0
    return df


def ff3(freq="m", vintage=None):
    """F-F 3팩터(소수). Mkt = Mkt-RF + RF(CRSP VW 총수익 — French 카드의 B)."""
    df = french("ff3_m" if freq == "m" else "ff3_d", "m" if freq == "m" else "d", vintage=vintage)
    df["Mkt"] = df["Mkt-RF"] + df["RF"]
    return df


def french_firm_mask(fid, cols=None, vintage=None, min_firms=MIN_FIRMS):
    """월 × 칸 → 기업 수 ≥ min_firms (F0: 미달 달은 «측정 불가»)."""
    n = french(fid, "nfirms", vintage=vintage, pct=False)
    cols = cols or LEGS.get(fid) or list(n.columns)
    return n[cols] >= min_firms


def french_vw_combine(fid, cols, vintage=None, size_lag=1, kind="vw_m"):
    """French 포트폴리오 몇 칸을 시가총액(기업 수 × 평균 규모, size_lag 달 전 행)으로 묶은 VW 수익(소수).
    T17 가치/성장 반쪽(BE-ME 십분위 6~10 · 1~5 — 설계 «t−1 시가총액 가중 · 파일 절 사용»). 한 칸이라도 없으면 NaN."""
    r = french(fid, kind, vintage=vintage)[cols]
    n = french(fid, "nfirms", vintage=vintage, pct=False)[cols]
    s = french(fid, "avgsize", vintage=vintage, pct=False)[cols]
    return vw_combine(r, n, s, size_lag=size_lag)


def vw_combine(r, n, s, size_lag=1):
    full = pd.period_range(r.index.min(), r.index.max(), freq="M")
    r, n, s = r.reindex(full), n.reindex(full), s.reindex(full)
    w = (n * s).shift(size_lag)
    ok = r.notna().all(axis=1) & w.notna().all(axis=1) & (w > 0).all(axis=1)
    out = (r * w).sum(axis=1) / w.sum(axis=1)
    return out.where(ok)


# ══════════════════════════════════════════════════════════════════════════
#  Shiller · FRED/ALFRED · Cboe · AQR · Wurgler · yfinance 파서
# ══════════════════════════════════════════════════════════════════════════
def shiller_month(v):
    """Shiller 날짜 1871.01 · 1871.1(=10월) · 1871.12 → Period."""
    y = int(v)
    m = int(round((float(v) - y) * 100))
    if not 1 <= m <= 12:
        raise ValueError("Shiller 날짜 %r" % v)
    return pd.Period("%04d-%02d" % (y, m), "M")


SHILLER_COLS = {"P": "P", "D": "D", "E": "E", "CPI": "CPI", "Rate GS10": "GS10_shiller"}   # CAPE 열은 읽지 않는다(설계 T13)


def parse_shiller_sheet(sh):
    """xlrd 시트(또는 nrows · ncols · cell_value 를 가진 것) → 월 DataFrame P · D · E · CPI · GS10_shiller."""
    hdr_r = None
    for r in range(min(sh.nrows, 40)):
        if str(sh.cell_value(r, 0)).strip() == "Date" and str(sh.cell_value(r, 1)).strip() == "P":
            hdr_r = r
            break
    if hdr_r is None:
        raise ValueError("Shiller 머리글(Date · P)을 못 찾았다")
    names = {str(sh.cell_value(hdr_r, c)).strip(): c for c in range(sh.ncols)}
    colmap = {out: names[src] for src, out in SHILLER_COLS.items() if src in names}
    rows, idx = [], []
    for r in range(hdr_r + 1, sh.nrows):
        v = sh.cell_value(r, 0)
        if not isinstance(v, float) or not 1800 < v < 2100:
            continue
        idx.append(shiller_month(v))
        rec = []
        for out, c in colmap.items():
            x = sh.cell_value(r, c)
            rec.append(float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else np.nan)
        rows.append(rec)
    df = pd.DataFrame(rows, index=pd.PeriodIndex(idx, freq="M"), columns=list(colmap))
    return df[~df.index.duplicated(keep="first")].sort_index()


def _xlrd():
    try:
        import xlrd
        return xlrd
    except ImportError:
        for p in (os.environ.get("TBATCH_PYLIB"), os.path.join(_cache_guard(), "pylib")):
            if p and os.path.isdir(p) and p not in sys.path:
                sys.path.insert(0, p)
        import xlrd                                     # noqa: F811 — 없으면 여기서 ImportError
        return xlrd


def parse_shiller_blob(blob):
    xlrd = _xlrd()
    wb = xlrd.open_workbook(file_contents=blob)
    return parse_shiller_sheet(wb.sheet_by_name("Data"))


def parse_fred_text(txt):
    """FRED/ALFRED 그래프 CSV → 월(PeriodIndex) 또는 일(DatetimeIndex) Series · 빈칸/«.» 는 NaN. 여러 열이면 DataFrame."""
    rows = list(csv.reader(io.StringIO(txt)))
    hdr = [h.strip() for h in rows[0]]
    body = [r for r in rows[1:] if r and r[0].strip()]
    dates = pd.to_datetime([r[0].strip() for r in body], format="%Y-%m-%d")
    vals = [[(float(x) if x.strip() not in ("", ".") else np.nan) for x in (r[1:] + [""] * (len(hdr) - len(r)))[:len(hdr) - 1]] for r in body]
    monthly = len(dates) > 1 and bool((dates.day == 1).all()) and 27 <= float(np.median(np.diff(dates.values).astype("timedelta64[D]").astype(float))) <= 32
    idx = pd.PeriodIndex(dates, freq="M") if monthly else pd.DatetimeIndex(dates)
    df = pd.DataFrame(vals, index=idx, columns=hdr[1:])
    df = df[~df.index.duplicated(keep="first")].sort_index()
    return df.iloc[:, 0].rename(hdr[1]) if df.shape[1] == 1 else df


def parse_cboe_text(txt):
    """Cboe 지수 이력 CSV(DATE MM/DD/YYYY) → 일 DataFrame(정렬 · 중복 첫 값)."""
    rows = list(csv.reader(io.StringIO(txt)))
    hdr = [h.strip().upper() for h in rows[0]]
    body = [r for r in rows[1:] if r and r[0].strip()]

    def iso(d):
        d = d.strip()
        if "/" in d:
            m, dd, y = d.split("/")
            return "%s-%02d-%02d" % (y, int(m), int(dd))
        return d
    idx = pd.DatetimeIndex(pd.to_datetime([iso(r[0]) for r in body], format="%Y-%m-%d"))
    vals = [[(float(x) if x.strip() not in ("", ".") else np.nan) for x in (r[1:] + [""] * len(hdr))[:len(hdr) - 1]] for r in body]
    df = pd.DataFrame(vals, index=idx, columns=hdr[1:])
    df.attrs["dup_dates"] = int(df.index.duplicated().sum())
    df.attrs["sorted_in_file"] = bool(df.index.is_monotonic_increasing)
    df = df[~df.index.duplicated(keep="first")].sort_index()
    return df


def _as_date(x):
    if isinstance(x, _dt.datetime):
        return x.date()
    if isinstance(x, _dt.date):
        return x
    if isinstance(x, str):
        s = x.strip()
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%Y%m%d"):
            try:
                return _dt.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


def parse_aqr_blob(blob, sheet):
    """AQR 자료 xlsx 한 시트 → 월 DataFrame(소수 그대로 · PeriodIndex). 머리글 = 첫 날짜 줄 바로 앞의 빈칸 아닌 줄."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        prev, hdr, idx, rows = None, None, [], []
        for r in ws.iter_rows(values_only=True):
            if not r:
                continue
            d = _as_date(r[0])
            if d is None:
                if any(x not in (None, "") for x in r):
                    prev = r
                continue
            if hdr is None:
                hdr = [("" if x is None else str(x).strip()) for x in (prev or [])]
            idx.append(pd.Period(d, "M"))
            rows.append([(float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else np.nan) for x in (list(r[1:]) + [None] * len(hdr))[:len(hdr) - 1]])
    finally:
        wb.close()
    cols = [c if c else "col%d" % j for j, c in enumerate(hdr[1:], 1)]
    df = pd.DataFrame(rows, index=pd.PeriodIndex(idx, freq="M"), columns=cols)
    df = df.loc[:, [c for c in df.columns if not c.startswith("col") or df[c].notna().any()]]
    return df[~df.index.duplicated(keep="first")].sort_index()


WURGLER_FORBIDDEN = ("recess",)                           # 경기후퇴 표지 — 신호로 쓰지 않는다(USREC 규칙과 같은 취지)


def parse_wurgler_blob(blob, keep_forbidden=False):
    """Wurgler SENTIMENT.xlsx DATA 시트 → 월 DataFrame(yearmo → Period). README «UPDATED» 는 attrs["updated"]."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        upd = None
        if "README" in wb.sheetnames:
            for r in wb["README"].iter_rows(values_only=True):
                s = " ".join(str(x) for x in r if x is not None)
                m = re.search(r"UPDATED:\s*([A-Za-z]+,?\s*\d{4})", s)
                if m:
                    upd = m.group(1)
                    break
        rows = list(wb["DATA"].iter_rows(values_only=True))
    finally:
        wb.close()
    hdr = [str(x).strip() if x is not None else "" for x in rows[0]]
    if hdr[0].lower() != "yearmo":
        raise ValueError("Wurgler DATA 머리글이 yearmo 로 시작하지 않는다: %r" % hdr[:3])
    idx, vals = [], []
    for r in rows[1:]:
        if r is None or r[0] is None:
            continue
        ym = str(int(r[0])) if isinstance(r[0], (int, float)) else str(r[0]).strip()
        if not re.fullmatch(r"\d{6}", ym):
            continue
        idx.append(pd.Period("%s-%s" % (ym[:4], ym[4:]), "M"))
        vals.append([(float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else np.nan) for x in (list(r[1:]) + [None] * len(hdr))[:len(hdr) - 1]])
    df = pd.DataFrame(vals, index=pd.PeriodIndex(idx, freq="M"), columns=hdr[1:])
    if not keep_forbidden:
        df = df.drop(columns=[c for c in WURGLER_FORBIDDEN if c in df.columns])
    df = df[~df.index.duplicated(keep="first")].sort_index()
    df.attrs["updated"] = upd
    return df


def parse_yf_blob(blob):
    df = pd.read_csv(io.BytesIO(blob), index_col=0)
    df.index = pd.DatetimeIndex(pd.to_datetime(df.index, format="%Y-%m-%d"))
    return df.sort_index()


# ── 로더(고정본만 읽는다) ─────────────────────────────────────────────────
_R = None


def _reg():
    global _R
    if _R is None:
        _R = registry()
    return _R


def shiller():
    """Shiller 월 자료 P(일 종가의 월평균) · D · E · CPI · GS10_shiller. CAPE 열은 주지 않는다(T13 은 재계산)."""
    return parse_shiller_blob(read_pinned(_reg()["shiller/ie_data"]))


def fred(series):
    ds = _reg()["fred/" + series]
    return parse_fred_text(read_pinned(ds).decode("utf-8"))


def cboe(name):
    return parse_cboe_text(read_pinned(_reg()["cboe/" + name]).decode("utf-8", "replace"))


def aqr(aid):
    ds = _reg()["aqr/" + aid]
    return parse_aqr_blob(read_pinned(ds), ds["sheet"])


def wurgler(keep_forbidden=False):
    return parse_wurgler_blob(read_pinned(_reg()["wurgler/SENTIMENT"]), keep_forbidden=keep_forbidden)


def yf_hist(ticker):
    return parse_yf_blob(read_pinned(_reg()["yf/" + ticker]))


def month_end(s):
    """일 계열 → 그달 마지막 관측(PeriodIndex)."""
    s = s.dropna()
    return s.groupby(s.index.to_period("M")).last()


def vix_month_end():
    return month_end(cboe("VIX")["CLOSE"])


def gspc_daily():
    return yf_hist("^GSPC")["Close"]


# ══════════════════════════════════════════════════════════════════════════
#  ALFRED 최초 공표 계열(T12)
# ══════════════════════════════════════════════════════════════════════════
def _month_end_dates(start, end_day):
    out, p = [], pd.Period(start, "M")
    while p.end_time.date() < end_day:
        out.append(p.end_time.strftime("%Y-%m-%d"))
        p += 1
    return out


def alfred_first_fetch(series="UNRATE", start=ALFRED_START, today=None, sleep=0.7):
    """월말마다의 판(그날 유효한 판)을 12개씩 받아 캐시에 둔다(ALFRED 그래프는 한 번에 열 12개까지)."""
    today = today or _dt.date.today()
    dates = _month_end_dates(start, today) + [today.strftime("%Y-%m-%d")]
    d = os.path.join(_cache_guard(), "raw", "alfred")
    os.makedirs(d, exist_ok=True)
    paths, failed = [], []

    def get(batch):
        fn = os.path.join(d, "%s_%s_%s.csv" % (series, batch[0], batch[-1]))
        if os.path.exists(fn):
            paths.append(fn)
            return
        url = "https://alfred.stlouisfed.org/graph/alfredgraph.csv?id=%s&vintage_date=%s" % (",".join([series] * len(batch)), ",".join(batch))
        try:
            blob, _ = _http_get(url, UA_PY, timeout=90)
        except urllib.error.HTTPError as e:
            if len(batch) > 1:
                time.sleep(sleep)
                for one in batch:
                    get([one])
                return
            failed.append((batch[0], e.code))
            return
        got = blob.decode("utf-8", "replace").split("\n", 1)[0].count(",")
        if got != len(batch):
            raise RuntimeError("ALFRED 가 열 %d개를 줬다(요청 %d) — 묶음 크기 한도가 바뀌었나" % (got, len(batch)))
        _write_bytes(fn, blob)
        paths.append(fn)
        time.sleep(sleep)

    for i in range(0, len(dates), 12):
        get(dates[i:i + 12])
    df = alfred_build_first([io.open(p, encoding="utf-8").read() for p in sorted(paths)], series)
    out = os.path.join(PUB, "%s_first_release.csv" % series)
    os.makedirs(PUB, exist_ok=True)
    txt = "obs_month,value,first_vintage,first_release\n" + "".join(
        "%s,%s,%s,%d\n" % (i, ("%.6g" % v.value), v.first_vintage, int(v.first_release)) for i, v in df.iterrows())
    io.open(out, "w", encoding="utf-8", newline="\n").write(txt)
    batch_sha = [(os.path.basename(p), sha256_bytes(open(p, "rb").read())) for p in sorted(paths)]
    ds = _reg()["alfred/UNRATE_first"]
    _write_meta(ds, {"key": ds["key"], "url": ds["url"], "fetched_at": _now(), "origin": "live (월말 판 표본)",
                     "n_batches": len(batch_sha), "batch_sha256": batch_sha, "failed_vintages": failed,
                     "sha256": sha256_bytes(txt.encode("utf-8")), "bytes": len(txt.encode("utf-8"))})
    return out


def alfred_build_first(texts, series="UNRATE"):
    """ALFRED 여러 판 CSV(열 = SERIES_YYYYMMDD) → 관측 달마다 «처음 실린 판»의 값.
    first_vintage = 그 값이 처음 보인 표본 판 날짜 · first_release = 판 날짜가 관측 달 + 3개월 안(표본 시작 전 이력은 거짓)."""
    vint = {}
    for t in texts:
        rows = list(csv.reader(io.StringIO(t)))
        hdr = [h.strip() for h in rows[0]]
        for j, h in enumerate(hdr[1:], 1):
            m = re.fullmatch(r"%s_(\d{8})" % re.escape(series), h)
            if not m:
                continue
            vd = "%s-%s-%s" % (m.group(1)[:4], m.group(1)[4:6], m.group(1)[6:])
            col = vint.setdefault(vd, {})
            for r in rows[1:]:
                if len(r) > j and r[j].strip() not in ("", "."):
                    col[r[0].strip()[:7]] = float(r[j])
    first = {}
    for vd in sorted(vint):
        for obs in sorted(vint[vd]):
            if obs not in first:
                first[obs] = (vint[vd][obs], vd)
    earliest = min(vint) if vint else None
    idx = pd.PeriodIndex(sorted(first), freq="M")
    val = [first[str(p)][0] for p in idx]
    fv = [first[str(p)][1] for p in idx]
    fr = [(pd.Period(v[:7], "M") - p).n <= 3 and not (v == earliest and (pd.Period(v[:7], "M") - p).n > 1) for p, v in zip(idx, fv)]
    return pd.DataFrame({"value": val, "first_vintage": fv, "first_release": fr}, index=idx)


def unrate_first_release():
    """저장소에 고정한 최초 공표 UNRATE(월) — value · first_vintage · first_release."""
    p = os.path.join(PUB, "UNRATE_first_release.csv")
    df = pd.read_csv(p, dtype={"obs_month": str, "first_vintage": str})
    df.index = pd.PeriodIndex(df.pop("obs_month"), freq="M")
    df["first_release"] = df["first_release"].astype(bool)
    return df


# ══════════════════════════════════════════════════════════════════════════
#  벤치마크 이음(B2 · final.tests.benchmarks)
# ══════════════════════════════════════════════════════════════════════════
def sp500_tr_splice(spx_daily, div_annual, tr_daily):
    """S&P 500 총수익 월 수익(소수) 이음 — 순수 함수.
    ^SP500TR 이 전월 말 값을 갖는 첫 달부터는 TR 월말 비, 그 앞(1975-02~)은 (SPX_t + D_t/12)/SPX_{t−1} − 1
    (Cboe SPX 월말 · Shiller D 연율 ÷ 12 — 근사, 등록 문서에 공개). 돌려주는 것 DataFrame(ret, src)."""
    spx = month_end(spx_daily)
    tr = month_end(tr_daily)
    first_full = tr.index.min() + 1                        # 1988-01-04 시작이면 1988-02 부터
    D = div_annual.reindex(spx.index)
    pre = (spx + D / 12.0) / spx.shift(1) - 1.0
    pre = pre[(pre.index < first_full)].dropna()
    post = (tr / tr.shift(1) - 1.0)
    post = post[post.index >= first_full].dropna()
    out = pd.concat([pd.DataFrame({"ret": pre, "src": "SPX+D/12"}), pd.DataFrame({"ret": post, "src": "SP500TR"})]).sort_index()
    return out[~out.index.duplicated(keep="last")]


def sp500_tr_monthly():
    return sp500_tr_splice(cboe("SPX")["SPX"], shiller()["D"], yf_hist("^SP500TR")["Close"])


def french_mkt(freq="m"):
    """French Mkt 총수익(소수) = Mkt-RF + RF."""
    return ff3(freq)["Mkt"]


# ══════════════════════════════════════════════════════════════════════════
#  가용 늦춤(final.tests.availability) — 선언된 것만 쓴다
# ══════════════════════════════════════════════════════════════════════════
#  이름 «출처:열» 또는 «출처». (재구성 recon, 실운용 live) 달 수. 월말 t 결정에는 «t − 늦춤» 달까지의 관측만 쓴다.
LAGS = {
    "french": (0, 2),          # 실시간 공개는 t−2 까지 → 장기 검정은 재구성(recon 0), 실운용은 S 층 대용
    "aqr": (0, 2),             # 약 6~7주 늦게 나온다
    "cboe": (0, 0), "yf": (0, 0),
    "fred:GS10": (1, 1), "fred:TB3MS": (1, 1), "fred:UNRATE": (1, 1), "fred:CPIAUCSL": (1, 1),
    "fred:DFII10": (0, 0),     # 일간 — 그날 종가(랩 D13 과 같은 가용)
    "alfred:UNRATE": (1, 1),   # + 최초 공표 판 날짜 ≤ 월말 t
    "shiller:P": (0, 0),       # 일 종가의 월평균 — 그달 말에 안다
    "shiller:E": (6, 6),       # 분기 E 선형 보간 · 1926 이전 연간 보간 → 6개월 늦춤(설계 T13)
    "shiller:CPI": (1, 1),     # 받은 파일에 공표 전 CPI 가 채워져 있다(비평1) → 1개월 늦춤
    "shiller:D": (6, 6),       # E 와 같은 보간 자료(분기 · 1926 이전 연간) → 6개월 늦춤(T13 tilt_M 의 1926-07 앞 이력 · t_signals.m12_excess)
    "shiller:GS10_shiller": (1, 1),   # 장기금리 월평균 — FRED GS10 과 같은 1개월(1926-07 앞 현금 대리)
    "wurgler": (3, 3),         # 심리 구성요소 3개월 늦춤(선언 · 설계 T15)
}
FORBIDDEN = {"fred:USREC": "USREC 는 신호로 쓰지 않는다(설계 availability)", "fred:USRECM": "USREC 계열",
             "wurgler:recess": "경기후퇴 표지 — USREC 와 같은 취지로 신호 금지"}
MODES = ("recon", "live")


def lag_months(name, mode="recon"):
    if mode not in MODES:
        raise ValueError("mode 는 recon · live")
    if name in FORBIDDEN:
        raise PermissionError("🚨 %s — %s" % (name, FORBIDDEN[name]))
    if name in LAGS:
        return LAGS[name][MODES.index(mode)]
    src = name.split(":")[0]
    if src in LAGS:
        return LAGS[src][MODES.index(mode)]
    raise KeyError("가용 늦춤이 선언되지 않은 입력: %s — LAGS 에 먼저 적는다(선언 없이 쓰지 않는다)" % name)


def horizon(name, t, mode="recon"):
    """월말 t 결정에 쓸 수 있는 마지막 관측 달."""
    return pd.Period(t, "M") - lag_months(name, mode)


def _avail_mask(obj, name, t, mode):
    h = horizon(name, t, mode)
    idx = obj.index
    if isinstance(idx, pd.PeriodIndex):
        m = np.asarray(idx <= h)
    elif isinstance(idx, pd.DatetimeIndex):
        m = np.asarray(idx <= h.end_time)
    else:                                          # 연간(정수 해) — 12월 말에야 그해가 끝난다
        last_year = h.year if h.month == 12 else h.year - 1
        m = np.asarray(pd.Index(idx).astype(int) <= last_year)
    if isinstance(obj, pd.DataFrame) and "first_vintage" in obj.columns:
        m = m & np.asarray(pd.to_datetime(obj["first_vintage"]) <= pd.Period(t, "M").end_time)
    return m


def asof(obj, name, t, mode="recon"):
    """obj 를 «월말 t 에 쓸 수 있는 데까지» 자른다(가용 늦춤 + ALFRED 판 날짜)."""
    return obj.loc[_avail_mask(obj, name, t, mode)]


def view_asof(inputs, t, mode="recon"):
    """inputs = {인자: (obj, 가용 이름)} → {인자: 잘린 obj}."""
    return {k: asof(o, nm, t, mode) for k, (o, nm) in inputs.items()}


def hold_next(sig):
    """월말 t 에 정한 상태 → t+1 보유(PeriodIndex + 1). 선견 점검은 상태에, 수익 정렬은 이것으로."""
    out = sig.copy()
    out.index = out.index + 1
    return out


# ══════════════════════════════════════════════════════════════════════════
#  선견 점검 틀(final.tests.lookahead_check · FACTORSWITCH 교훈)
# ══════════════════════════════════════════════════════════════════════════
def _row(x, t):
    try:
        v = x.loc[t]
    except KeyError:
        return None
    return np.atleast_1d(np.asarray(v, dtype=float))


def _same(a, b, atol):
    if a is None or b is None or a.shape != b.shape:
        return False
    na, nb = np.isnan(a), np.isnan(b)
    return bool((na == nb).all() and np.allclose(a[~na], b[~nb], atol=atol, rtol=0))


def _poison(obj, name, t, mode, rng):
    avail = _avail_mask(obj, name, t, mode)
    bad = ~avail
    if not bad.any():
        return obj
    out = obj.copy()
    if isinstance(out, pd.Series):
        out = out.astype(float)
        out.iloc[np.flatnonzero(bad)] = rng.normal(0.0, 37.0, int(bad.sum()))
        return out
    for c in out.columns:
        if c in ("first_vintage", "first_release") or not pd.api.types.is_numeric_dtype(out[c]):
            continue
        col = out[c].astype(float).to_numpy().copy()
        col[bad] = rng.normal(0.0, 37.0, int(bad.sum()))
        out[c] = col
    return out


def lookahead_check(signal_fn, inputs, n=LOOKAHEAD_N, seed=LOOKAHEAD_SEED, t_min=None, t_max=None, mode="recon",
                    poison=True, atol=1e-10):
    """상태는 t 에서 정하고 보유는 t+1 — 무작위 t n개에서 입력을 «t 에 쓸 수 있는 데까지» 잘라 다시 계산한 신호가
    원 신호와 같아야 한다(절단 시험). poison=True 면 가용 지평 뒤 값을 난수로 바꾸고(색인은 그대로) 한 번 더 —
    위치 색인 · shift(−k) · 전 표본 적률 누수를 잡는다.
      signal_fn(**{인자: obj}) → PeriodIndex(월말 t) Series/DataFrame
      inputs = {인자: (obj, 가용 이름)}  — 가용 이름은 LAGS 에 선언된 것
    돌려주는 것 {"ok", "n", "n_bad", "bad"(≤5 · t · 원 값 · 잘린 값), "mode", "seed"} — 수익은 다루지 않는다."""
    full = signal_fn(**{k: o for k, (o, nm) in inputs.items()})
    fr = full.to_frame() if isinstance(full, pd.Series) else full
    cand = fr.index[fr.notna().all(axis=1).to_numpy()]
    if t_min is not None:
        cand = cand[cand >= pd.Period(t_min, "M")]
    if t_max is not None:
        cand = cand[cand <= pd.Period(t_max, "M")]
    if len(cand) == 0:
        return {"ok": False, "n": 0, "n_bad": 0, "bad": [], "mode": mode, "seed": seed, "err": "원 신호에 유효한 t 가 없다"}
    rng = np.random.default_rng(seed)
    pick = np.sort(rng.choice(len(cand), size=min(n, len(cand)), replace=False))
    prng = np.random.default_rng(seed + 1)
    bad = []
    for i in pick:
        t = cand[i]
        want = _row(fr, t)
        tests = [("절단", view_asof(inputs, t, mode))]
        if poison:
            tests.append(("독", {k: _poison(o, nm, t, mode, prng) for k, (o, nm) in inputs.items()}))
        for how, args in tests:
            try:
                got = signal_fn(**args)
                got = _row(got.to_frame() if isinstance(got, pd.Series) else got, t)
                err = None
            except Exception as e:
                got, err = None, repr(e)[:200]
            if not _same(want, got, atol):
                bad.append({"t": str(t), "test": how, "full": None if want is None else want.tolist(),
                            "recomputed": None if got is None else got.tolist(), "err": err})
    return {"ok": not bad, "n": int(len(pick)), "n_bad": len(bad), "bad": bad[:5], "mode": mode, "seed": seed}


def require_no_lookahead(res, label=""):
    """틀리면 실행하지 않는다(설계 문구)."""
    if not res.get("ok"):
        raise SystemExit("🚨 선견 점검 실패 %s — %d/%d · 첫 사례 %s · 실행하지 않는다." % (label, res.get("n_bad", 0), res.get("n", 0),
                                                                          (res.get("bad") or [{}])[0].get("t")))
    return True


# ══════════════════════════════════════════════════════════════════════════
#  명세(data/_tb_manifest.json) — 행 수 · 기간 · 해시만
# ══════════════════════════════════════════════════════════════════════════
def _rng(idx):
    if len(idx) == 0:
        return None, None
    a, b = idx.min(), idx.max()
    f = (lambda x: x.strftime("%Y-%m-%d")) if isinstance(idx, pd.DatetimeIndex) else str
    return f(a), f(b)


def _col_span(s):
    v = s.dropna()
    a, b = _rng(v.index)
    inner = 0
    if len(v):
        seg = s.loc[v.index.min():v.index.max()]
        inner = int(seg.isna().sum())
    return {"first": a, "last": b, "n": int(len(v)), "inner_missing": inner}


def summarize(ds, blob):
    """고정본 하나의 구조 요약 — 수익 통계 없음."""
    k = ds["kind"]
    out = {}
    if k == "french":
        f = parse_french_blob(blob)
        main = next(x for x in ("vw_m", "m", "d") if x in f["blocks"])
        blocks = []
        for kind, df in f["blocks"].items():
            a, b = _rng(df.index)
            blocks.append({"kind": kind, "title": f["titles"][kind], "rows": int(len(df)), "first": a, "last": b,
                           "cols": f["cols"][kind], "n_missing_codes": f["n_missing"][kind], "dup_dates": f["dup_dates"][kind]})
        a, b = _rng(f["blocks"][main].index)
        out.update({"inner": f["inner"], "crsp_vintage": f["crsp"], "main_block": main, "rows": int(len(f["blocks"][main])),
                    "first": a, "last": b, "blocks": blocks, "preamble_head": f["preamble"][:2]})
        legs = LEGS.get(ds["fid"])
        if legs and "nfirms" in f["blocks"]:
            nf = f["blocks"]["nfirms"]
            vw = f["blocks"][main]
            lq = {}
            for c in legs:
                if c not in nf.columns:
                    lq[c] = "칸 없음"
                    continue
                low = nf[c] < MIN_FIRMS
                lq[c] = {"months_lt_%d" % MIN_FIRMS: int(low.sum()), "first_ok": _rng(nf.index[~low.to_numpy()])[0],
                         "vw_missing": int(vw[c].isna().sum()) if c in vw.columns else None}
            out["legs_firms"] = lq
        if ds["fid"] == "ff3_m" and "m" in f["blocks"]:
            m = f["blocks"]["m"]["Mkt-RF"]
            hit = sum(1 for p, v in SPOT_MKTRF.items() if pd.Period(p, "M") in m.index and abs(m.loc[pd.Period(p, "M")] - v) < 0.0051)
            present = sum(1 for p in SPOT_MKTRF if pd.Period(p, "M") in m.index)
            out["spot_mktrf"] = "%d/%d 일치(설계 스폿 %d개 중 이 판에 있는 달 %d)" % (hit, present, len(SPOT_MKTRF), present)
    elif k == "shiller":
        df = parse_shiller_blob(blob)
        a, b = _rng(df.index)
        out.update({"rows": int(len(df)), "first": a, "last": b, "cols": {c: _col_span(df[c]) for c in df.columns}})
    elif k == "fred":
        s = parse_fred_text(blob.decode("utf-8"))
        s = s if isinstance(s, pd.Series) else s.iloc[:, 0]
        a, b = _rng(s.index)
        out.update({"rows": int(len(s)), "first": a, "last": b, "n_missing": int(s.isna().sum()),
                    "freq": "M" if isinstance(s.index, pd.PeriodIndex) else "D", "last_valid": _rng(s.dropna().index)[1]})
    elif k == "cboe":
        df = parse_cboe_text(blob.decode("utf-8", "replace"))
        a, b = _rng(df.index)
        out.update({"rows": int(len(df)), "first": a, "last": b, "cols": list(df.columns), "dup_dates": df.attrs.get("dup_dates"),
                    "sorted_in_file": df.attrs.get("sorted_in_file"), "n_missing": int(df.isna().sum().sum())})
    elif k == "aqr":
        df = parse_aqr_blob(blob, ds["sheet"])
        a, b = _rng(df.index)
        use = [c for c in ("TSMOM", "USA") if c in df.columns]
        out.update({"sheet": ds["sheet"], "rows": int(len(df)), "first": a, "last": b, "n_cols": int(df.shape[1]),
                    "used_cols": {c: _col_span(df[c]) for c in use}})
    elif k == "wurgler":
        df = parse_wurgler_blob(blob)
        a, b = _rng(df.index)
        comp = [c for c in ("SENT", "SENT_ORTH", "pdnd", "ripo", "nipo", "cefd", "s") if c in df.columns]
        out.update({"rows": int(len(df)), "first": a, "last": b, "readme_updated": df.attrs.get("updated"),
                    "cols": {c: _col_span(df[c]) for c in comp}})
    elif k == "yf":
        df = parse_yf_blob(blob)
        a, b = _rng(df.index)
        dv = df["Dividends"] if "Dividends" in df.columns else pd.Series(dtype=float)
        sp = df["Stock Splits"] if "Stock Splits" in df.columns else pd.Series(dtype=float)
        out.update({"rows": int(len(df)), "first": a, "last": b, "cols": list(df.columns), "n_dividends": int((dv > 0).sum()),
                    "n_splits": int((sp > 0).sum()), "nan_close": int(df["Close"].isna().sum())})
    elif k == "alfred_first":
        df = pd.read_csv(io.BytesIO(blob), dtype={"obs_month": str, "first_vintage": str})
        fr = df["first_release"].astype(bool)
        delay = [(pd.Period(v[:7], "M") - pd.Period(o, "M")).n for o, v in zip(df["obs_month"], df["first_vintage"])]
        dly = pd.Series(delay)[fr.to_numpy()]
        out.update({"rows": int(len(df)), "first": df["obs_month"].iloc[0], "last": df["obs_month"].iloc[-1],
                    "first_release_from": df["obs_month"][fr].iloc[0] if fr.any() else None, "n_not_first_release": int((~fr).sum()),
                    "earliest_vintage": df["first_vintage"].min(), "latest_vintage": df["first_vintage"].max(),
                    "release_delay_months": {str(k2): int(v) for k2, v in dly.value_counts().sort_index().items()}})
    return out


def _expect_diff(ds, summ):
    e = ds.get("expect")
    if not e:
        return None
    rows, first, last = (list(e) + [None, None, None])[:3]
    d = {}
    if rows is not None and summ.get("rows") != rows:
        d["rows"] = [rows, summ.get("rows")]
    if first is not None and summ.get("first") and not str(summ["first"]).startswith(first):
        d["first"] = [first, summ.get("first")]
    if last is not None and summ.get("last") and not str(summ["last"]).startswith(last):
        d["last"] = [last, summ.get("last")]
    return d or "설계 기대와 같다"


def design_state(pinned_sha, design_sha):
    if design_sha is None:
        return "설계 해시 없음"
    if pinned_sha is None:
        return "고정본 없음"
    return "일치" if pinned_sha == design_sha else "다름"


def _cmp_frames(a, b):
    """두 판의 같은 표 — 행 수 · 더해진/빠진 색인 · 겹친 칸 중 바뀐 칸 수 · 최대 |차|(원 단위). 수익 통계가 아니라 개정 셈이다."""
    if isinstance(a, pd.Series):
        a = a.to_frame()
    if isinstance(b, pd.Series):
        b = b.to_frame()
    common = a.index.intersection(b.index)
    cols = [c for c in a.columns if c in b.columns]
    A = a.loc[common, cols].to_numpy(float)
    B = b.loc[common, cols].to_numpy(float)
    na, nb = np.isnan(A), np.isnan(B)
    diff = np.where(na | nb, 0.0, np.abs(A - B))
    changed = (diff > 1e-9) | (na ^ nb)
    added = b.index.difference(a.index)
    removed = a.index.difference(b.index)
    return {"rows": [int(len(a)), int(len(b))], "added": int(len(added)), "added_range": list(_rng(added)) if len(added) else None,
            "removed": int(len(removed)), "changed_cells": int(changed.sum()), "cells_compared": int(changed.size),
            "max_abs_diff": round(float(diff.max()), 6) if diff.size else 0.0,
            "cols_only_pinned": [c for c in a.columns if c not in b.columns], "cols_only_live": [c for c in b.columns if c not in a.columns]}


def compare_blobs(ds, pinned, live):
    """고정본 대 지금 판 — 해시가 다를 때만 부른다."""
    k = ds["kind"]
    if k == "french":
        fa, fb = parse_french_blob(pinned), parse_french_blob(live)
        out = {"crsp_vintage": [fa["crsp"], fb["crsp"]]}
        for kind in fa["blocks"]:
            if kind in fb["blocks"] and kind in ("vw_m", "m", "d", "ew_m", "nfirms", "avgsize"):
                out[kind] = _cmp_frames(fa["blocks"][kind], fb["blocks"][kind])
        return out
    if k == "fred":
        return _cmp_frames(parse_fred_text(pinned.decode("utf-8")), parse_fred_text(live.decode("utf-8")))
    if k == "cboe":
        return _cmp_frames(parse_cboe_text(pinned.decode("utf-8", "replace")), parse_cboe_text(live.decode("utf-8", "replace")))
    if k == "aqr":
        return _cmp_frames(parse_aqr_blob(pinned, ds["sheet"]), parse_aqr_blob(live, ds["sheet"]))
    if k == "wurgler":
        return _cmp_frames(parse_wurgler_blob(pinned), parse_wurgler_blob(live))
    if k == "shiller":
        return _cmp_frames(parse_shiller_blob(pinned), parse_shiller_blob(live))
    return None


def build_manifest(R=None, live=None, write=True):
    R = R or registry()
    old = {}
    if os.path.exists(MANIFEST):
        try:
            old = json.load(io.open(MANIFEST, encoding="utf-8")).get("datasets", {})
        except Exception:
            old = {}
    D = {}
    for key, ds in R.items():
        p = pinned_path(ds)
        rec = {"source": ds["src"], "url": ds["url"], "license": ds["license"], "cards": ds["cards"],
               "pinned_in": ("repo:data/_tb_fred/" if ds["src"] in PUBLIC_SRC else "cache:raw/%s/" % ds["src"]) + ds["fname"],
               "design_sha256": ds["design_sha"]}
        if not os.path.exists(p):
            rec.update({"state": "고정본 없음", "design_state": design_state(None, ds["design_sha"])})
            D[key] = rec
            continue
        with open(p, "rb") as f:
            blob = f.read()
        sha = sha256_bytes(blob)
        meta = _read_meta(ds) or {}
        if meta.get("sha256") != sha:                     # 캐시 메타가 없거나 낡았으면 옛 명세에서(같은 해시일 때만)
            o = old.get(key, {}).get("pinned", {})
            meta = o if o.get("sha256") == sha else {"origin": "알 수 없음(메타 없음)", "sha256": sha}
        pin = {"sha256": sha, "bytes": len(blob), "fetched_at": meta.get("fetched_at"), "last_modified": meta.get("last_modified"),
               "origin": meta.get("origin"), "fetched_from": meta.get("url")}
        for x in ("n_batches", "batch_sha256", "failed_vintages", "yfinance"):
            if x in meta:
                pin[x] = meta[x]
        rec["pinned"] = pin
        rec["design_state"] = design_state(sha, ds["design_sha"])
        try:
            rec["summary"] = summarize(ds, blob)
            rec["vs_design_expect"] = _expect_diff(ds, rec["summary"])
        except Exception as e:
            rec["summary"] = {"error": repr(e)[:300]}
        lv = (live or {}).get(key)
        if lv is None and old.get(key, {}).get("live", {}).get("pinned_sha256") == sha:
            rec["live"] = old[key]["live"]
        elif lv is not None:
            if "error" in lv:
                rec["live"] = {"checked_at": lv.get("fetched_at"), "error": lv["error"]}
            else:
                same = lv["sha256"] == sha
                rec["live"] = {"checked_at": lv["fetched_at"], "sha256": lv["sha256"], "last_modified": lv.get("last_modified"),
                               "same_as_pinned": same, "pinned_sha256": sha, "matches_design": (lv["sha256"] == ds["design_sha"]) if ds["design_sha"] else None}
                if not same:
                    try:
                        with open(lv["path"], "rb") as f:
                            rec["live"]["diff"] = compare_blobs(ds, blob, f.read())
                    except Exception as e:
                        rec["live"]["diff"] = {"error": repr(e)[:300]}
        D[key] = rec
    doc = {"note": ("배치 T 자료 판 동결(B0) — build/t_data.py 가 쓴다. 값은 없다: URL · 받은 때 · SHA-256 · 행 수 · 기간 · 구조만. "
                    "AQR · Cboe · Wurgler(D3) · French · Shiller · yfinance 원자료는 저장소 밖 캐시에만 있고, "
                    "연방 통계(FRED/ALFRED)만 data/_tb_fred/ 에 고정한다. 사이트 자료가 아니다(D1 — L 층 전용)."),
           "generated_at": _now(), "cache_root": "$TBATCH_CACHE 또는 %TEMP%/tbatch_cache (저장소 밖)",
           "design": "tbatch_research.json final · data.hashes / verified_this_step(설계 워크플로 2026-09-25~26)",
           "decisions": {"D1": "장기 검정은 L-R · L-C 기전 증거 층 전용 · 사이트 10년 규약 유지",
                         "D2": "S&P 지수 선물 ±5% · 지수 옵션 · 커버드콜 ETF 허용 / 교차자산 선물 불허 → T02 측정만",
                         "D3": "AQR · Cboe · Wurgler 원자료와 파생 NAV 비커밋 · 비게시 · 해시만",
                         "D4": "배치 FWER 0.025 → H_T0 하나 · 누적 N 은 DSR 보고만"},
           "availability": {k: {"recon": v[0], "live": v[1]} for k, v in LAGS.items()},
           "forbidden_signals": FORBIDDEN,
           "cache_identity": {"file": CACHE_ID.replace(os.sep, "/"), "sha256": cache_id_sha(),
                              "note": "러너는 이 표지를 가진 캐시에서만 굽는다(무작위 값 · 뜻 없음 — 캐시 뿌리를 바꿔 다시 굽는 길을 막는다)"},
           "pylib": pylib_pins(),
           "datasets": D}
    for k in ("from_cache_pylib",):
        doc["pylib"][PYLIB_PKG].pop(k, None)                   # 어디서 불렀는지는 PC 마다 다르다 — 파일 SHA 만 고정한다
    if write:
        os.makedirs(os.path.dirname(MANIFEST), exist_ok=True)
        io.open(MANIFEST, "w", encoding="utf-8", newline="\n").write(json.dumps(doc, ensure_ascii=False, indent=1) + "\n")
    return doc


def check_pins(keys=None):
    """고정본 SHA == 명세 — 실행 전 관문. 돌려주는 것 (ok, 문제 목록)."""
    doc = json.load(io.open(MANIFEST, encoding="utf-8"))
    R = registry()
    bad = []
    for key in keys or R:
        rec = doc["datasets"].get(key)
        if rec is None or "pinned" not in rec:
            bad.append("%s: 명세에 고정본 없음" % key)
            continue
        p = pinned_path(R[key])
        if not os.path.exists(p):
            bad.append("%s: 고정본 파일 없음" % key)
            continue
        with open(p, "rb") as f:
            sha = sha256_bytes(f.read())
        if sha != rec["pinned"]["sha256"]:
            bad.append("%s: SHA 다름(%s… ≠ %s…)" % (key, sha[:12], rec["pinned"]["sha256"][:12]))
    return not bad, bad


def report(doc):
    """범위 · 판 보고(값 없음)."""
    lines = []
    for key, r in doc["datasets"].items():
        s = r.get("summary") or {}
        lv = r.get("live") or {}
        live = ("같음" if lv.get("same_as_pinned") else ("바뀜" if lv.get("same_as_pinned") is False else ("오류" if lv.get("error") else "—")))
        exp = r.get("vs_design_expect")
        lines.append("%-30s %-11s %6s  %-10s → %-10s  설계해시 %-8s 지금판 %-4s %s%s" % (
            key, str((r.get("pinned") or {}).get("origin") or "없음")[:11], s.get("rows", "—"), s.get("first", "—"), s.get("last", "—"),
            r.get("design_state"), live, ("" if exp in (None, "설계 기대와 같다") else "기대차 %s " % json.dumps(exp, ensure_ascii=False)),
            ("CRSP %s" % s["crsp_vintage"]) if s.get("crsp_vintage") else ""))
    return "\n".join(lines)


def b0_markdown(doc=None):
    """등록 문서 B0 절 — 판 동결 표(해시 전체 · Last-Modified · 받은 때 · 행 · 기간)와 French 절 머리글."""
    doc = doc or json.load(io.open(MANIFEST, encoding="utf-8"))
    L = ["## B0 판 동결 — 자료 해시(data/_tb_manifest.json 과 같다)", "",
         "| 키 | 출처 URL | Last-Modified | 받은 때(UTC) | SHA-256 | 행 | 기간 | 설계 해시 |", "|---|---|---|---|---|---|---|---|"]
    for key, r in doc["datasets"].items():
        p = r.get("pinned") or {}
        s = r.get("summary") or {}
        L.append("| %s | %s | %s | %s | `%s` | %s | %s ~ %s | %s |" % (key, r["url"], p.get("last_modified") or "—", p.get("fetched_at") or "—",
                                                                  p.get("sha256", "—"), s.get("rows", "—"), s.get("first", "—"), s.get("last", "—"),
                                                                  r.get("design_state")))
    L += ["", "### French 절 머리글(주 절 · 판 표식)", ""]
    for key, r in doc["datasets"].items():
        s = r.get("summary") or {}
        if r["source"].startswith("french") and s.get("blocks"):
            L.append("- **%s** (`%s` · CRSP %s)" % (key, s.get("inner"), s.get("crsp_vintage")))
            for b in s["blocks"]:
                L.append("  - %s «%s» %d행 %s~%s · 열: %s" % (b["kind"], b["title"], b["rows"], b["first"], b["last"], ", ".join(b["cols"])))
    return "\n".join(L) + "\n"


# ══════════════════════════════════════════════════════════════════════════
#  눈가린 연기 시험(실자료 · 값은 안 본다)
# ══════════════════════════════════════════════════════════════════════════
def smoke():
    import qbatch_core as Q                                  # blind_smoke — 표준출력 버림 · 모양만 돌려준다

    def lk():
        f = ff3("m")
        d = ff3("d")

        def sig(mkt, day):                                   # 합성 틀 점검용 상태(T01 SLOW 꼴 · 판정에 쓰지 않는다)
            slow = mkt["Mkt-RF"].rolling(12).mean().ge(0).astype(float).where(mkt["Mkt-RF"].rolling(12).count() == 12)
            s2 = day["Mkt-RF"].rolling(126).var()
            v = s2.groupby(s2.index.to_period("M")).last().reindex(mkt.index)
            return pd.DataFrame({"slow": slow, "var_hi": (v > v.expanding(36).median()).astype(float).where(v.expanding(36).count() >= 36)})
        res = lookahead_check(sig, {"mkt": (f, "french:ff3_m"), "day": (d, "french:ff3_d")}, n=40)
        if not res["ok"]:
            raise RuntimeError("선견 틀 실자료 점검 실패 %d/%d" % (res["n_bad"], res["n"]))
        return res

    tests = {
        "french_vw_combine(beme 가치 반쪽)": (french_vw_combine, ("beme", BEME_VALUE)),
        "french_vw_combine(beme 성장 반쪽)": (french_vw_combine, ("beme", BEME_GROWTH)),
        "french_firm_mask(전 다리)": (lambda: {f: french_firm_mask(f) for f in LEGS}, ()),
        "sp500_tr_monthly": (sp500_tr_monthly, ()),
        "unrate_first_release": (unrate_first_release, ()),
        "vix_month_end": (vix_month_end, ()),
        "wurgler(구성요소)": (wurgler, ()),
        "aqr TSMOM · BAB": (lambda: (aqr("TSMOM"), aqr("BAB")), ()),
        "shiller": (shiller, ()),
        "lookahead_check(French 월 · 일 · 40개)": (lk, ()),
    }
    out = {}
    for nm, (fn, a) in tests.items():
        r = Q.blind_smoke(fn, *a)
        out[nm] = {"ok": r["ok"], "sec": r["sec"], "shape": r["shape"] if r["ok"] else None, "err": (r["err"] or "")[-300:] or None}
    return out


# ══════════════════════════════════════════════════════════════════════════
#  합성 자료 시험
# ══════════════════════════════════════════════════════════════════════════
def _zip_text(name, txt):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr(name, txt.encode("latin-1"))
    return b.getvalue()


def _st_french():
    hdr = ",<= 0,Lo 30,Hi 30,Lo 10,2-Dec,Hi 10\n"
    txt = ("This file was created using the 202607 CRSP database.\nIt contains value- and equal-weighted returns.\n\n"
           "Missing data are indicated by -99.99 or -999.\n\n\n"
           "  Value Weight Returns -- Monthly\n" + hdr +
           "192607,   1.00,   2.00, -99.99,   3.00,   4.00,   5.00\n192608,  -1.00,  -2.00,   0.50,  -3.00,  -4.00,  -5.00\n"
           "192609,   0.10,   0.20,   0.30,   0.40,   0.50,   0.60\n\n"
           "  Equal Weight Returns -- Monthly\n" + hdr +
           "192607,   1.10,   2.10,   3.10,   3.20,   4.20,   5.20\n192608,  -1.10,  -2.10,   0.60,  -3.10,  -4.10,  -5.10\n"
           "192609,   0.11,   0.21,   0.31,   0.41,   0.51,   0.61\n\n"
           "  Value Weight Returns -- Annual from January to December\n" + hdr +
           "1926,   5.00,   6.00,   7.00,   8.00,   9.00,  10.00\n1927,  -5.00,  -6.00,  -7.00,  -8.00,  -9.00, -10.00\n\n"
           "  Number of Firms in Portfolios\n" + hdr +
           "192607,    10,    25,    30,    19,    22,    40\n192608,    10,    25,    30,    20,    22,    40\n"
           "192609,    10,    25,    30,    21,    22,    40\n\n"
           "  Average Firm Size\n" + hdr +
           "192607, -99.99,  100.0,  200.0,   50.0,   60.0,  300.0\n192608,  -999,  110.0,  210.0,   55.0,   65.0,  310.0\n"
           "192609,   1.0,  120.0,  220.0,   57.0,   66.0,  320.0\n\n"
           "  Sum of BE / Sum of ME\n" + hdr + "1926,   0.50,   0.60,   0.70,   0.80,   0.90,   1.00\n\n"
           "  Value Weight Average of BE / ME\n" + hdr + "1926,   0.51,   0.61,   0.71,   0.81,   0.91,   1.01\n\n"
           "Copyright 2026 Eugene F. Fama and Kenneth R. French\n")
    f = parse_french_blob(_zip_text("Portfolios_Formed_on_BE-ME.csv", txt))
    B = f["blocks"]
    assert set(B) == {"vw_m", "ew_m", "vw_a", "nfirms", "avgsize", "sum_ratio", "vwavg_a"}, sorted(B)
    assert f["crsp"] == "202607" and f["inner"] == "Portfolios_Formed_on_BE-ME.csv"
    assert list(B["vw_m"].columns) == ["<= 0", "Lo 30", "Hi 30", "Lo 10", "Dec 2", "Hi 10"]
    assert B["vw_m"].shape == (3, 6) and isinstance(B["vw_m"].index, pd.PeriodIndex)
    assert np.isnan(B["vw_m"].loc[pd.Period("1926-07", "M"), "Hi 30"]) and f["n_missing"]["vw_m"] == 1
    assert f["n_missing"]["avgsize"] == 2 and np.isnan(B["avgsize"].iloc[1, 0])
    assert list(B["vw_a"].index) == [1926, 1927]
    assert B["nfirms"].loc[pd.Period("1926-07", "M"), "Lo 10"] == 19
    # FF3 꼴(제목 없는 첫 절 + 연간) · 일간 · 옛 TXT
    ff = ("This file was created using the 202607 CRSP database.\nThe 1-month TBill ...\n\n,Mkt-RF,SMB,HML,RF\n"
          "192607,   2.89,  -2.42,  -2.75,   0.22\n192608,   2.64,  -1.44,   4.13,   0.25\n\n Annual Factors: January-December \n"
          ",Mkt-RF,SMB,HML,RF\n1927,  29.47,  -2.04,  -3.49,   3.12\n\nCopyright 2026\n")
    g = parse_french_blob(_zip_text("F-F_Research_Data_Factors.csv", ff))
    assert set(g["blocks"]) == {"m", "a"} and g["titles"]["m"] == "" and g["blocks"]["a"].index[0] == 1927
    dd = "created by using the 202607 CRSP database.\n\n,Mkt-RF,SMB,HML,RF\n19260701,    0.09,   -0.25,   -0.27,    0.01\n19260702,    0.45,   -0.33,   -0.06,    0.01\n"
    h = parse_french_blob(_zip_text("F-F_Research_Data_Factors_daily.csv", dd))
    assert set(h["blocks"]) == {"d"} and isinstance(h["blocks"]["d"].index, pd.DatetimeIndex)
    tx = ("This file was created by CMPT_ME_BEME_RETS using the 200507 CRSP database.\nThe 1-month TBill return is from Ibbotson.\n\n"
          "        Mkt-RF     SMB     HML      RF\n192607    2.65   -2.38   -3.03    0.22\n192608    2.52   -1.28    4.35    0.25\n\n"
          " Annual Factors: January-December\n        Mkt-RF     SMB     HML      RF\n1927   29.47   -2.04   -3.49    3.12\n")
    k = parse_french_blob(_zip_text("Historical_Archives/08 2005 Update/ftp/F-F_Research_Data_Factors.txt", tx))
    assert k["crsp"] == "200507" and list(k["blocks"]["m"].columns) == ["Mkt-RF", "SMB", "HML", "RF"] and k["inner"] == "F-F_Research_Data_Factors.txt"
    assert abs(k["blocks"]["m"].iloc[1, 2] - 4.35) < 1e-12 and set(k["blocks"]) == {"m", "a"}
    # 머리글과 자료 사이 빈 줄(현재 VAR 판 VW 연간 절) — 제목 · 칸 이름을 제대로 잡는다 · 칸 수가 안 맞는 줄은 머리글로 안 쓴다
    vh = ",Lo 20,Hi 20\n"
    vt = ("This file was created using the 202607 CRSP database.\n\n  Value Weighted Returns -- Monthly\n" + vh +
          "196307,   1.00,   2.00\n196308,   1.10,   2.10\n\n  Value Weighted Returns -- Annual from January to December\n" + vh + "\n"
          "1964,   5.00,   6.00\n1965,   5.10,   6.10\n\n  Some Title Without Header\n\n196307,   1.0,   2.0,   3.0\n")
    q = parse_french_blob(_zip_text("Portfolios_Formed_on_VAR.csv", vt))
    assert set(q["blocks"]) == {"vw_m", "vw_a", "some_title_without_header"} or set(q["blocks"]) == {"vw_m", "vw_a", "m"}, sorted(q["blocks"])
    assert list(q["blocks"]["vw_a"].columns) == ["Lo 20", "Hi 20"] and q["titles"]["vw_a"].startswith("Value Weighted Returns -- Annual")
    odd = [b for b in q["blocks"] if b not in ("vw_m", "vw_a")][0]
    assert list(q["cols"][odd]) == [] or q["blocks"][odd].shape[1] == 3
    # 기업 수 마스크 · VW 묶기(t−1 가중)
    r = B["vw_m"][["Lo 10", "Dec 2"]] / 100
    n, s = B["nfirms"][["Lo 10", "Dec 2"]], B["avgsize"][["Lo 10", "Dec 2"]]
    v = vw_combine(r, n, s, size_lag=1)
    w8 = np.array([19 * 50.0, 22 * 60.0])
    want8 = float((np.array([-0.03, -0.04]) * w8).sum() / w8.sum())
    assert np.isnan(v.iloc[0]) and abs(v.iloc[1] - want8) < 1e-15
    s2 = s.copy()
    s2.iloc[1, 1] = np.nan
    assert np.isnan(vw_combine(r, n, s2).iloc[2])
    assert list((B["nfirms"]["Lo 10"] >= MIN_FIRMS).to_numpy()) == [False, True, True]
    return "French 다절 · FF3 · 일간 · TXT · 결측 부호 · 열 이름 정규화 · VW 묶기 · 기업 수 마스크"


class _FakeSheet:
    def __init__(self, rows):
        self.rows = rows
        self.nrows = len(rows)
        self.ncols = max(len(r) for r in rows)

    def cell_value(self, r, c):
        row = self.rows[r]
        return row[c] if c < len(row) else ""


def _st_shiller():
    assert str(shiller_month(1871.1)) == "1871-10" and str(shiller_month(1871.01)) == "1871-01"
    assert str(shiller_month(2026.09)) == "2026-09" and str(shiller_month(1999.12)) == "1999-12"
    rows = [["Stock Market Data"], ["", "S&P"], ["Date", "P", "D", "E", "CPI", "Fraction", "Rate GS10", "Real Price", "x", "y", "z", "w", "CAPE"],
            [1871.01, 4.44, 0.26, 0.4, 12.46, 1871.04, 5.32, 118.9, 0, 0, 0, 0, "NA"],
            [1871.1, 4.5, 0.26, 0.4, 12.8, 1871.8, 5.3, 119.0, 0, 0, 0, 0, "NA"],
            [2026.09, 6600.0, "", "", 333.89, 2026.7, 4.1, 1, 0, 0, 0, 0, 41.0], ["", "", ""]]
    df = parse_shiller_sheet(_FakeSheet(rows))
    assert list(df.columns) == ["P", "D", "E", "CPI", "GS10_shiller"] and "CAPE" not in df.columns
    assert [str(p) for p in df.index] == ["1871-01", "1871-10", "2026-09"] and np.isnan(df.loc[pd.Period("2026-09", "M"), "E"])
    return "Shiller 날짜(1871.1 = 10월) · 머리글 찾기 · NA · CAPE 열 안 읽음"


def _st_text_parsers():
    s = parse_fred_text("observation_date,GS10\n1953-04-01,2.83\n1953-05-01,.\n1953-06-01,3.11\n")
    assert isinstance(s.index, pd.PeriodIndex) and s.name == "GS10" and np.isnan(s.iloc[1]) and len(s) == 3
    d = parse_fred_text("observation_date,DFII10\n2003-01-02,2.43\n2003-01-03,2.43\n2003-01-06,\n")
    assert isinstance(d.index, pd.DatetimeIndex) and np.isnan(d.iloc[2])
    c = parse_cboe_text("DATE,OPEN,HIGH,LOW,CLOSE\n01/03/1990,1,2,0.5,1.5\n01/02/1990,17.24,17.24,17.24,17.24\n01/03/1990,9,9,9,9\n")
    assert list(c.columns) == ["OPEN", "HIGH", "LOW", "CLOSE"] and c.attrs["dup_dates"] == 1 and not c.attrs["sorted_in_file"]
    assert str(c.index[0].date()) == "1990-01-02" and c["CLOSE"].iloc[1] == 1.5 and len(c) == 2
    y = parse_yf_blob(b"Date,Open,High,Low,Close,Adj Close,Volume,Dividends,Stock Splits\n2020-01-02,1,1,1,1,0.9,10,0,0\n2020-01-03,2,2,2,2,1.8,10,0.1,0\n")
    assert len(y) == 2 and y["Dividends"].iloc[1] == 0.1
    return "FRED 월/일 · Cboe(MM/DD/YYYY · 중복 · 정렬) · yfinance CSV"


def _st_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TSMOM Factors"
    ws.append(["AQR Capital Management, LLC — Time Series Momentum: Factors, Monthly"])
    ws.append(["This file contains the excess returns of the long/short TSMOM factors."])
    ws.append([None, "TSMOM", "TSMOM^CM", "TSMOM^EQ"])
    ws.append([_dt.datetime(1985, 1, 31), 0.01, 0.02, None])
    ws.append([_dt.datetime(1985, 2, 28), -0.01, "n/a", 0.03])
    ws2 = wb.create_sheet("BAB Factors")
    ws2.append(["notes"])
    ws2.append(["DATE", "AUS", "USA"])
    ws2.append(["01/31/1931", 0.1, 0.2])
    b = io.BytesIO()
    wb.save(b)
    a = parse_aqr_blob(b.getvalue(), "TSMOM Factors")
    assert list(a.columns) == ["TSMOM", "TSMOM^CM", "TSMOM^EQ"] and str(a.index[1]) == "1985-02" and np.isnan(a.iloc[1, 1])
    bb = parse_aqr_blob(b.getvalue(), "BAB Factors")
    assert list(bb.columns) == ["AUS", "USA"] and str(bb.index[0]) == "1931-01"
    wb = openpyxl.Workbook()
    rd = wb.active
    rd.title = "README"
    rd.append(["INVESTOR SENTIMENT DATA"])
    rd.append(["UPDATED: March, 2026"])
    dt = wb.create_sheet("DATA")
    dt.append(["yearmo", "SENT", "SENT_ORTH", "pdnd", "ripo", "nipo", "cefd", "s", "recess"])
    dt.append([196507, None, None, 1.0, None, 3.0, 4.0, 5.0, 0])
    dt.append([196508, 0.1, 0.2, 1.1, 2.1, 3.1, 4.1, 5.1, 1])
    b = io.BytesIO()
    wb.save(b)
    w = parse_wurgler_blob(b.getvalue())
    assert "recess" not in w.columns and w.attrs["updated"] == "March, 2026" and str(w.index[0]) == "1965-07"
    assert "recess" in parse_wurgler_blob(b.getvalue(), keep_forbidden=True).columns
    return "AQR xlsx(머리글 · 날짜 · 비수치) · Wurgler DATA/README · recess 기본 제외"


def _st_alfred():
    v1 = "observation_date,UNRATE_19991231,UNRATE_20000131\n1999-10-01,4.1,4.1\n1999-11-01,4.1,4.2\n1999-12-01,,4.0\n"
    v2 = ("observation_date,UNRATE_20000229,UNRATE_20000331,UNRATE_20000430\n1999-10-01,4.1,4.1,4.1\n1999-11-01,4.2,4.2,4.3\n"
          "1999-12-01,4.0,4.0,4.0\n2000-01-01,,4.1,4.1\n2000-02-01,,,4.2\n")
    df = alfred_build_first([v2, v1])
    assert [str(p) for p in df.index] == ["1999-10", "1999-11", "1999-12", "2000-01", "2000-02"]
    assert df.loc[pd.Period("1999-11", "M"), "value"] == 4.1                 # 개정(4.2 · 4.3)을 따르지 않는다
    assert df.loc[pd.Period("1999-12", "M"), "first_vintage"] == "2000-01-31"
    assert df.loc[pd.Period("2000-01", "M"), "first_vintage"] == "2000-03-31"  # 늦게 나온 달(한 판 건너뜀)
    assert not df.loc[pd.Period("1999-10", "M"), "first_release"]            # 표본 시작 전 이력(가장 이른 판 · 2개월 전)
    assert df.loc[pd.Period("1999-11", "M"), "first_release"]
    # 가용: 2000-02 말에는 2000-01 값이 아직 없다(판 날짜 2000-03-31)
    a = asof(df, "alfred:UNRATE", "2000-02")
    assert str(a.index.max()) == "1999-12"
    a = asof(df, "alfred:UNRATE", "2000-03")
    assert str(a.index.max()) == "2000-01"
    return "ALFRED 최초 공표(개정 무시 · 늦은 공표 · 표본 전 이력 표지 · 판 날짜 가용)"


def _st_avail():
    s = pd.Series(np.arange(24.0), index=pd.period_range("2000-01", periods=24, freq="M"))
    assert str(asof(s, "fred:UNRATE", "2000-12").index.max()) == "2000-11"
    assert str(asof(s, "shiller:E", "2000-12").index.max()) == "2000-06"
    assert str(asof(s, "shiller:CPI", "2000-12").index.max()) == "2000-11"
    assert str(asof(s, "wurgler:pdnd", "2000-12").index.max()) == "2000-09"
    assert str(asof(s, "french:beme", "2000-12").index.max()) == "2000-12"
    assert str(asof(s, "french:beme", "2000-12", mode="live").index.max()) == "2000-10"
    assert str(asof(s, "aqr:TSMOM", "2000-12", mode="live").index.max()) == "2000-10"
    dd = pd.Series(1.0, index=pd.date_range("2000-01-01", "2000-03-31", freq="D"))
    assert asof(dd, "cboe:VIX", "2000-02").index.max() == pd.Timestamp("2000-02-29")
    an = pd.Series([1.0, 2.0, 3.0], index=pd.Index([1998, 1999, 2000]))
    assert asof(an, "french:beme", "2000-11").index.max() == 1999 and asof(an, "french:beme", "2000-12").index.max() == 2000
    for bad, exc in (("fred:USREC", PermissionError), ("wurgler:recess", PermissionError), ("fred:BAA", KeyError), ("shiller:CAPE", KeyError)):
        try:
            lag_months(bad)
            raise AssertionError("%s 가 통과했다" % bad)
        except exc:
            pass
    h = hold_next(pd.Series([1.0], index=pd.PeriodIndex(["2000-01"], freq="M")))
    assert str(h.index[0]) == "2000-02"
    return "늦춤 표(recon · live) · 월/일/연 자르기 · 금지(USREC · recess) · 미선언 거부 · t+1 보유"


def _st_lookahead():
    rng = np.random.default_rng(7)
    idx = pd.period_range("1960-01", "2005-12", freq="M")
    r = pd.DataFrame({"Mkt-RF": rng.normal(0.005, 0.045, len(idx))}, index=idx)
    u = pd.Series(5 + np.cumsum(rng.normal(0, 0.1, len(idx))), index=idx, name="UNRATE")

    def clean(mkt, un):
        slow = mkt["Mkt-RF"].rolling(12).mean()
        fast = mkt["Mkt-RF"]
        w = 0.5 * np.sign(slow) + 0.5 * np.sign(fast)
        un_m = un.reindex(mkt.index)                         # un 은 t−1 까지만 보인다 → t−1 값과 t−12..t−1 평균
        prev = un_m.shift(1)
        up = (prev > prev.rolling(12).mean()).astype(float).where(prev.rolling(12).count() == 12)
        med = w.expanding(36).median()                       # 확장창 적률 — 허용
        return pd.DataFrame({"w": w, "up": up, "w_vs_med": (w > med).astype(float).where(w.expanding(36).count() >= 36)})

    inp = {"mkt": (r, "french:ff3_m"), "un": (u, "fred:UNRATE")}
    ok = lookahead_check(clean, inp, n=60)
    assert ok["ok"] and ok["n"] == 60, ok

    def leak_future(mkt, un):
        return mkt["Mkt-RF"].shift(-1).gt(0).astype(float).where(mkt["Mkt-RF"].shift(-1).notna())

    def leak_avail(mkt, un):                                 # UNRATE_t 를 쓴다(가용 t−1 위반)
        un_m = un.reindex(mkt.index)
        return (un_m > un_m.rolling(12).mean()).astype(float).where(un_m.rolling(12).count() == 12)

    def leak_fullsample(mkt, un):                            # 전 표본 중앙값으로 중심화
        x = mkt["Mkt-RF"].rolling(12).mean()
        return x - x.median()

    def leak_position(mkt, un):                              # 위치 색인(np.roll) — 절단에선 NaN 이 되지만 독 시험이 잡는다
        v = np.roll(mkt["Mkt-RF"].to_numpy(), -1)
        return pd.Series(v > 0, index=mkt.index).astype(float)

    for nm, fn in (("미래 shift", leak_future), ("가용 위반", leak_avail), ("전 표본 적률", leak_fullsample), ("위치 색인", leak_position)):
        res = lookahead_check(fn, inp, n=60)
        assert not res["ok"] and res["n_bad"] > 0, (nm, res)
    try:
        require_no_lookahead(lookahead_check(leak_future, inp, n=10), "합성")
        raise AssertionError("require_no_lookahead 가 멈추지 않았다")
    except SystemExit:
        pass
    return "선견 틀: 깨끗한 신호 통과 · 미래 shift · 가용 위반 · 전 표본 적률 · 위치 색인 넷 모두 잡음 · 실패 시 멈춤"


def _st_splice():
    days = pd.bdate_range("1987-10-01", "1988-03-31")
    spx = pd.Series(np.linspace(300, 330, len(days)), index=days)
    D = pd.Series(12.0, index=pd.period_range("1987-01", "1988-12", freq="M"))
    trd = days[days >= "1988-01-04"]
    tr = pd.Series(np.linspace(1000, 1100, len(trd)), index=trd)
    out = sp500_tr_splice(spx, D, tr)
    assert [str(p) for p in out.index] == ["1987-11", "1987-12", "1988-01", "1988-02", "1988-03"]
    assert list(out["src"]) == ["SPX+D/12"] * 3 + ["SP500TR"] * 2
    me = month_end(spx)
    want = (me[pd.Period("1987-12", "M")] + 1.0) / me[pd.Period("1987-11", "M")] - 1
    assert abs(out.loc[pd.Period("1987-12", "M"), "ret"] - want) < 1e-15
    tme = month_end(tr)
    assert abs(out.loc[pd.Period("1988-02", "M"), "ret"] - (tme.iloc[1] / tme.iloc[0] - 1)) < 1e-15
    return "S&P TR 이음(SPX + D/12 → ^SP500TR 첫 완전 월)"


def _st_manifest():
    assert sha256_bytes(b"abc") == "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
    assert design_state("x", "x") == "일치" and design_state("x", "y") == "다름" and design_state("x", None) == "설계 해시 없음"
    a = pd.DataFrame({"x": [1.0, 2.0, np.nan]}, index=pd.period_range("2000-01", periods=3, freq="M"))
    b = pd.DataFrame({"x": [1.0, 2.5, np.nan, 4.0]}, index=pd.period_range("2000-01", periods=4, freq="M"))
    c = _cmp_frames(a, b)
    assert c["added"] == 1 and c["changed_cells"] == 1 and c["max_abs_diff"] == 0.5 and c["removed"] == 0
    R = registry()
    assert len([k for k in R if k.startswith("french/")]) == 13 and "alfred/UNRATE_first" in R and "yf/^GSPC" in R
    assert all(R[k]["design_sha"] for k in R if k.startswith("french/"))
    need = {"french/ff3_m", "french/ff3_d", "french/mom_d", "french/prior12_2", "french/beta", "french/var", "french/op", "french/beme",
            "french/ep", "french/cfp", "french/dp", "french/ni", "french/ind12", "shiller/ie_data", "fred/GS10", "fred/TB3MS", "fred/UNRATE",
            "fred/CPIAUCSL", "alfred/UNRATE_first", "cboe/BXM", "cboe/BXMD", "cboe/PUT", "cboe/VIX", "cboe/SPX", "aqr/TSMOM", "aqr/BAB",
            "wurgler/SENTIMENT", "yf/^GSPC", "yf/^SP500TR", "yf/SPY", "yf/BIL"}
    assert need <= set(R), sorted(need - set(R))
    for k, ds in R.items():                                  # 사용 허락 없는 출처는 저장소에 고정하지 않는다
        assert (ds["src"] in PUBLIC_SRC) == pinned_path(ds).startswith(PUB), k
    return "해시 · 설계 대조 · 판 비교 · 목록(카드가 쓰는 자료 전부 · French 설계 해시) · 공개 경계"


def _st_static():
    import ast
    src = io.open(os.path.abspath(__file__), encoding="utf-8").read()
    bad = []
    for nd in ast.walk(ast.parse(src)):
        if isinstance(nd, ast.Call) and isinstance(nd.func, ast.Name) and nd.func.id == "open":
            if any(kw.arg == "encoding" for kw in nd.keywords):
                continue
            mode = nd.args[1].value if len(nd.args) > 1 and isinstance(nd.args[1], ast.Constant) else ""
            if "b" not in str(mode):
                bad.append(nd.lineno)
    assert not bad, "open() encoding 없음: %s" % bad
    return "open() encoding(validate_site 규칙) 정적 점검"


def _st_guard():
    global CACHE
    old = CACHE
    try:
        CACHE = os.path.join(ROOT, "data", "_x_cache")
        try:
            _cache_guard()
            raise AssertionError("저장소 안 캐시가 통과했다")
        except SystemExit:
            pass
    finally:
        CACHE = old
    # 캐시 표지(selftest 는 임시 캐시) · 백업 경계
    assert cache_id_sha() is None and cache_id_init() and not cache_id_init()
    s = cache_id_sha()
    assert check_cache_identity({"cache_identity": {"sha256": s}})[0]
    assert not check_cache_identity({"cache_identity": {"sha256": "0" * 64}})[0] and not check_cache_identity({})[0]
    try:
        backup(os.path.join(ROOT, "data", "_x_backup"))
        raise AssertionError("저장소 안 백업이 통과했다")
    except SystemExit:
        pass
    return "캐시가 저장소 안이면 멈춤(D3) · 캐시 표지(명세와 같아야) · 저장소 안 백업 거부"


def selftest():
    global CACHE, PUB, MANIFEST, _R
    saved = (CACHE, PUB, MANIFEST, _R)
    res, ok = [], True
    with tempfile.TemporaryDirectory() as td:
        CACHE, PUB, MANIFEST, _R = os.path.join(td, "cache"), os.path.join(td, "pub"), os.path.join(td, "m.json"), None
        try:
            for fn in (_st_french, _st_shiller, _st_text_parsers, _st_xlsx, _st_alfred, _st_avail, _st_lookahead, _st_splice,
                       _st_manifest, _st_static, _st_guard):
                try:
                    res.append(("통과", fn.__name__, fn()))
                except Exception as e:
                    ok = False
                    import traceback
                    res.append(("실패", fn.__name__, traceback.format_exc()[-800:]))
        finally:
            CACHE, PUB, MANIFEST, _R = saved
    for st, nm, msg in res:
        print("  %s %-18s %s" % ("✓" if st == "통과" else "✗", nm, msg))
    print("selftest %d/%d 통과" % (sum(1 for r in res if r[0] == "통과"), len(res)))
    return 0 if ok else 1


# ══════════════════════════════════════════════════════════════════════════
def main(argv):
    if "--selftest" in argv:
        return selftest()
    R = registry()
    if "--seed" in argv:
        d = argv[argv.index("--seed") + 1]
        for k, v in seed(d, R).items():
            print("  %-30s %s" % (k, v))
    if "--fetch" in argv:
        for k, ds in R.items():
            try:
                p = fetch(ds)
                print("  %-30s %s" % (k, "고정본 " + os.path.basename(p)))
            except Exception as e:
                print("  %-30s 🚨 받기 실패 %s" % (k, repr(e)[:160]))
    live = None
    if "--live" in argv:
        live = {}
        for k, ds in R.items():
            lv = fetch_live(ds)
            if lv is not None:
                live[k] = lv
                print("  %-30s 지금 판 %s" % (k, lv.get("sha256", "")[:12] or lv.get("error", "")[:80]))
    if "--manifest" in argv or live is not None:
        doc = build_manifest(R, live=live)
        print(report(doc))
        print("→ %s" % os.path.relpath(MANIFEST, ROOT))
    if "--cache-id-init" in argv:
        print("캐시 표지: %s" % ("새로 만들었다" if cache_id_init() else "이미 있다"))
    if "--check" in argv:
        ok, bad = check_pins()
        print("고정본 점검: %s" % ("통과" if ok else "실패 %d" % len(bad)))
        for b in bad:
            print("  " + b)
        ci, cmsg = check_cache_identity()
        pl, pmsg = check_pylib()
        print("캐시 표지: %s · pylib: %s" % (cmsg, pmsg))
        if not (ok and ci and pl):
            return 1
    if "--backup" in argv:
        i = argv.index("--backup")
        d = argv[i + 1] if i + 1 < len(argv) and not argv[i + 1].startswith("--") else None
        r = backup(d)
        print("백업: 복사 %d · 이미 같음 %d · 캐시에 없음 %d%s" % (r["copied"], r["same"], len(r["missing"]),
                                                            " · ⚠ %TEMP% 안이다" if r["dest_in_temp"] else ""))
    if "--restore" in argv:
        i = argv.index("--restore")
        d = argv[i + 1] if i + 1 < len(argv) and not argv[i + 1].startswith("--") else None
        print("복원: %d개" % restore(d)["restored"])
    if "--smoke" in argv:
        for k, v in smoke().items():
            print("  %s %-40s %5ss %s" % ("✓" if v["ok"] else "✗", k, v["sec"], "" if v["ok"] else v["err"]))
    if "--b0" in argv:
        i = argv.index("--b0")
        md = b0_markdown()
        if i + 1 < len(argv) and not argv[i + 1].startswith("--"):
            io.open(argv[i + 1], "w", encoding="utf-8", newline="\n").write(md)
            print("→ %s" % argv[i + 1])
        else:
            print(md)
    if len(argv) <= 1:
        print(__doc__)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
